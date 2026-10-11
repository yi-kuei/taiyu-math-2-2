# -*- coding: utf-8 -*-
"""
generate_lesson_plan_docx.py
將《2-2_三角函數的定義與性質_四節課教學計畫與教案.md》
編譯為具備微軟原生 OMML 數學方程式的正式 A4 Word 教案文件 (.docx)
"""

import os
import re
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree
import latex2mathml.converter

# 載入微軟 MML2OMML.XSL
XSL_PATHS = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
]
XSL_FILE = next((p for p in XSL_PATHS if os.path.exists(p)), None)
TRANSFORM = etree.XSLT(etree.parse(XSL_FILE)) if XSL_FILE else None

def latex_to_omml(latex_str, color_hex=None):
    if not TRANSFORM:
        return None
    try:
        clean = latex_str.replace(r'\dfrac', r'\frac').replace(r'\,', ' ').replace(r'\approx', '~').replace(r'\cdot', r'\times')
        mml = latex2mathml.converter.convert(clean)
        omml = TRANSFORM(etree.fromstring(mml.encode('utf-8')))
        xml = re.sub(r'<\?xml[^>]+\?>', '', etree.tostring(omml, encoding='utf-8').decode('utf-8')).strip()
        if color_hex:
            xml = re.sub(r'(<m:r>|<m:r\s+[^>]*>)', rf'\1<m:rPr><w:rPr><w:color w:val="{color_hex}"/></w:rPr></m:rPr>', xml)
        return parse_xml(xml)
    except:
        return None

def add_math_runs(paragraph, full_text, default_color=None, bold=False, font_size=Pt(10.5)):
    parts = re.split(r'(\$[^\$]+\$)', full_text)
    for part in parts:
        if not part:
            continue
        if part.startswith('$') and part.endswith('$') and len(part) > 2:
            latex_expr = part[1:-1].strip()
            color_hex = f"{default_color[0]:02X}{default_color[1]:02X}{default_color[2]:02X}" if default_color else None
            omml_node = latex_to_omml(latex_expr, color_hex=color_hex)
            if omml_node is not None:
                paragraph._p.append(omml_node)
            else:
                r = paragraph.add_run(latex_expr)
                r.font.name = "Times New Roman"
                r.font.size = font_size
                r.bold = bold
                if default_color:
                    r.font.color.rgb = RGBColor(*default_color)
        else:
            r = paragraph.add_run(part)
            r.font.name = "微軟正黑體"
            r.font.size = font_size
            r.bold = bold
            if default_color:
                r.font.color.rgb = RGBColor(*default_color)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color_hex="CBD5E1"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color_hex}"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def make_lesson_plan_doc():
    doc = docx.Document()
    
    # 版面設定：A4 直向 (210mm x 297mm)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    
    # 頁首
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("泰宇版 技術高中數學（第一冊）2-2 三角函數的定義與性質 ‧ 四節課課程計畫與教案")
    hrun.font.name = "微軟正黑體"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(140, 140, 140)
    
    # 頁尾
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("PLVS 數學科教學研究會 ‧ 名偵探柯南數理探案特訓 旗艦全功能標準版")
    frun.font.name = "微軟正黑體"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(140, 140, 140)

    # 1. 主標題
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r1 = p_title.add_run("泰宇版 技術高中數學（第一冊）單元 2-2\n三角函數的定義與性質 課程計畫與教學教案")
    r1.font.name = "微軟正黑體"
    r1.font.size = Pt(17)
    r1.bold = True
    r1.font.color.rgb = RGBColor(120, 53, 15) # 琥珀深棕

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r2 = p_sub.add_run("【名偵探柯南數理探案特訓 ‧ 旗艦全功能教學標準版】（4 節課 × 50 分鐘，共 200 分鐘）")
    r2.font.name = "微軟正黑體"
    r2.font.size = Pt(11)
    r2.bold = True
    r2.font.color.rgb = RGBColor(180, 83, 9)

    # 2. 基本資料表
    info_tbl = doc.add_table(rows=4, cols=4)
    info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_tbl, "D97706")
    
    info_data = [
        [("設計單位", True), ("PLVS 數學科教學研究會", False), ("實施年級", True), ("技術型高中一年級", False)],
        [("教材版本", True), ("泰宇版 數學（第一冊）2-2 (P.107~124)", False), ("教學總時數", True), ("4 節課（共 200 分鐘）", False)],
        [("教學資源", True), ("單檔 HTML 互動平台、Word 原生 OMML 雙欄學習單、60張圖解卡 PDF", False), ("授權密碼", True), ("88888888 (教師詳解)", False)],
        [("核心素養", True), ("數-U-A1 系統思考與符號運用、數-U-B1 科技資訊與幾何測距、數-U-C2 生活工程坡度素養", False), ("評量工具", True), ("60道自檢題、60秒速算PK、JIS B4 雙欄學習單", False)],
    ]
    
    col_widths = [Mm(25), Mm(62), Mm(25), Mm(62)]
    for r_idx, row in enumerate(info_data):
        for c_idx, (text, is_hdr) in enumerate(row):
            cell = info_tbl.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if is_hdr:
                set_cell_shading(cell, "FEF3C7")
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                add_math_runs(cp, text, default_color=(120, 53, 15), bold=True, font_size=Pt(9.5))
            else:
                set_cell_shading(cell, "FFFFFF")
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                add_math_runs(cp, text, default_color=(51, 65, 85), bold=False, font_size=Pt(9.5))
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. 壹、總體課程計畫架構與時數分配總表
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    hr1 = h1.add_run("壹、總體課程計畫架構與時數分配總表")
    hr1.font.name = "微軟正黑體"
    hr1.font.size = Pt(13)
    hr1.bold = True
    hr1.font.color.rgb = RGBColor(30, 58, 138)

    plan_tbl = doc.add_table(rows=5, cols=5)
    plan_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(plan_tbl, "CBD5E1")
    
    plan_headers = ["節次", "核心主題", "課本範圍", "搭配圖解卡與工具", "核心學習目標與教學任務"]
    plan_widths = [Mm(18), Mm(32), Mm(26), Mm(34), Mm(64)]
    
    for c_idx, h_text in enumerate(plan_headers):
        cell = plan_tbl.cell(0, c_idx)
        cell.width = plan_widths[c_idx]
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        set_cell_shading(cell, "1E3A8A")
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(0)
        r = cp.add_run(h_text)
        r.font.name = "微軟正黑體"
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    plan_rows = [
        (
            "第一節\n(50 min)",
            "銳角三角比定義與\n雙重視角切換實戰",
            "課本 P.107~112\n例題 1~3\n隨堂 1~3",
            "CASE 01~11\n🔬 直角測距儀\n📝 學習單模組一&二",
            "1. 透過相似三角形建立「三角函數是比值非長度」概念。\n2. 熟記 $\\sin, \\cos, \\tan, \\cot, \\sec, \\csc$ 定義與「S配C、T配T」倒數伴侶。\n3. 直角三角形 $\\angle A$ 與 $\\angle B$ 視角對邊鄰邊換位實戰。\n4. 規範「設比例參數 $k$」搭配畢氏定理求邊長與各函數值。"
        ),
        (
            "第二節\n(50 min)",
            "特別角推導、生活科技\n坡度與 60 秒速算特訓",
            "課本 P.113~116\n例題 4~5\n隨堂 4~5",
            "CASE 12~20\n⚡ 60秒速算挑戰\n📝 學習單範例4",
            "1. 幾何推導正三角形對折與等腰直角之邊長比。\n2. 熟練度數與弳度量混合運算（如 $\\sin\\frac{\\pi}{6}$）。\n3. 融入營建署無障礙坡道生活素養，用計算機按 $\\tan^{-1}$。\n4. 班級 60 秒速算 PK，建立特別角直覺反射神經。"
        ),
        (
            "第三節\n(50 min)",
            "神級魔法六角密碼盤\n與四大恆等關係神技",
            "課本 P.117~121\n例題 6~8\n隨堂 6~8",
            "CASE 21~35\n🔯 魔法六角密碼盤\n📝 學習單重點3&範例8",
            "1. 掌握六角盤「左無co、右帶co、中為1」佈局。\n2. 熟練對角互倒、倒三角平方、橫向餘角、相鄰商數。\n3. 靈活運用互餘角轉換將非特別角相消。\n4. 突破常規：一次正餘弦齊次分式「同除以 $\\cos\\theta$」秒殺技巧。"
        ),
        (
            "第四節\n(50 min)",
            "三角恆等式深度破案、\n是非題辨析與劇場版考核",
            "課本 P.121~124\n例題 9、習題是非\n基礎題、進階題",
            "CASE 36~60\n🏆 少年偵探團考核\n📝 學習單模組三&四",
            "1. 掌握和差連環扣 $(\\sin\\pm\\cos)^2=1\\pm 2\\sin\\cos$ 與 $\\tan+\\cot=\\frac{1}{\\sin\\cos}$。\n2. 深度診斷課本是非題 5 大迷思（比值非長度、同角平方等）。\n3. 攻克劇場版立方和公式進階題，完成 60 案晉級名偵探！"
        )
    ]
    
    for r_idx, r_data in enumerate(plan_rows, start=1):
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, text in enumerate(r_data):
            cell = plan_tbl.cell(r_idx, c_idx)
            cell.width = plan_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_shading(cell, bg_col)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            if c_idx == 0:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_math_runs(cp, text, default_color=(120, 53, 15), bold=True, font_size=Pt(9))
            elif c_idx in (1, 2, 3):
                add_math_runs(cp, text, default_color=(30, 41, 59), bold=(c_idx==1), font_size=Pt(9))
            else:
                add_math_runs(cp, text, default_color=(51, 65, 85), bold=False, font_size=Pt(8.5))

    doc.add_page_break()

    # 4. 貳、分節詳細教學教案
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    hr2 = h2.add_run("貳、分節詳細教學教案 (每節 50 分鐘)")
    hr2.font.name = "微軟正黑體"
    hr2.font.size = Pt(13)
    hr2.bold = True
    hr2.font.color.rgb = RGBColor(30, 58, 138)

    def add_lesson_section(lesson_num, lesson_title, lesson_desc, objectives, timeline_rows):
        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(10)
        lp.paragraph_format.space_after = Pt(2)
        r = lp.add_run(f"【第 {lesson_num} 節課教案】{lesson_title}")
        r.font.name = "微軟正黑體"
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(180, 83, 9)
        
        # 目標條列
        op = doc.add_paragraph()
        op.paragraph_format.space_after = Pt(4)
        add_math_runs(op, "教學目標： " + " ｜ ".join(objectives), default_color=(71, 85, 105), bold=False, font_size=Pt(9))
        
        # 教學活動表
        act_tbl = doc.add_table(rows=len(timeline_rows)+1, cols=5)
        act_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(act_tbl, "CBD5E1")
        act_widths = [Mm(18), Mm(22), Mm(56), Mm(52), Mm(26)]
        
        act_hdrs = ["時間", "教學環節", "教師教學活動 (Teacher's Activities)", "學生學習活動 (Students' Activities)", "教學資源與工具"]
        for c_idx, h_text in enumerate(act_hdrs):
            cell = act_tbl.cell(0, c_idx)
            cell.width = act_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            set_cell_shading(cell, "D97706")
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(0)
            ar = cp.add_run(h_text)
            ar.font.name = "微軟正黑體"
            ar.font.size = Pt(9)
            ar.bold = True
            ar.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, r_data in enumerate(timeline_rows, start=1):
            bg = "FFFBEB" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, text in enumerate(r_data):
                cell = act_tbl.cell(r_idx, c_idx)
                cell.width = act_widths[c_idx]
                set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
                set_cell_shading(cell, bg)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                if c_idx == 0:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_math_runs(cp, text, default_color=(180, 83, 9), bold=True, font_size=Pt(8.5))
                elif c_idx == 1:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_math_runs(cp, text, default_color=(30, 58, 138), bold=True, font_size=Pt(8.5))
                else:
                    add_math_runs(cp, text, default_color=(51, 65, 85), bold=False, font_size=Pt(8.5))
                    
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 4.1 第一節課
    add_lesson_section(
        1,
        "銳角三角函數的定義與視角切換實戰 (50 分鐘)",
        "聚焦直角三角比幾何本質、六大函數符號與雙重視角切換",
        [
            "理解三角比是邊長比值而非長度",
            "熟記六大三角函數定義與倒數伴侶",
            "掌握 $\\angle A$ 與 $\\angle B$ 視角對邊鄰邊切換",
            "熟練設比例參數 $k$ 搭配畢氏定理求邊長"
        ],
        [
            ("00~08 分\n(8 min)", "引起動機\n概念奠基", "1. 投影大螢幕開啟【阿笠博士動態測距儀】，拖曳角度滑桿，引導學生觀察三角形形變中「同角度比值完全守恆」本質。\n2. 呈現 CASE 01：回顧相似形與正三角形對折，揭開「三角比」序幕。", "1. 觀察大螢幕幾何形變與數值連動。\n2. 歸納發現：只要角度固定，直角三角形無論大小，對邊/斜邊比值完全相同！", "阿笠博士測距儀\nCASE 01 圖解卡\n觸控大螢幕"),
            ("08~20 分\n(12 min)", "核心概念\n六大函數", "1. 投影 CASE 02 & CASE 03，講解六個函數符號與口訣：\n   · 正弦 $\\sin\\theta=\\frac{y}{r}$、餘割 $\\csc\\theta=\\frac{r}{y}$（S配C）\n   · 餘弦 $\\cos\\theta=\\frac{x}{r}$、正割 $\\sec\\theta=\\frac{r}{x}$（C配S）\n   · 正切 $\\tan\\theta=\\frac{y}{x}$、餘切 $\\cot\\theta=\\frac{x}{y}$（T配T）\n2. 發放 JIS B4 雙欄學習單，指導填寫【模組一：重點 1】。", "1. 大聲朗誦口訣：「S配C，T配T；首字必互換！」。\n2. 於學習單【重點 1】正確填入符號與比值定義。", "CASE 02、03\nWord 學習單\n（模組一）"),
            ("20~35 分\n(15 min)", "範例隨堂\n雙重視角", "1. 投影 CASE 04 & CASE 06（8-15-17 三角形）：\n   · 站在 $\\angle A$ 看：對邊是 15，鄰邊是 8。\n   · 站在 $\\angle B$ 看：對邊是 8，鄰邊是 15！\n   · 翻轉卡片展示「名偵探推理筆記」三步驟線索。\n2. 學生自主完成 CASE 05 & 07（5-12-13 三角形）。", "1. 練習「換位思考」：手指指出目標角，對面是對邊，腳下是鄰邊。\n2. 完成學習單【模組二：CASE 04 & 06】實戰練習。", "CASE 04~07\n白板註記畫筆\nWord 學習單"),
            ("35~45 分\n(10 min)", "知一求其餘\n參數法突破", "1. 投影 CASE 08（$\\cos A=\\frac{5}{13}$，求其餘五者）：\n   · 破解迷思：$\\cos A=\\frac{5}{13}$ 不代表鄰邊一定是 5、斜邊一定是 13！\n   · 規範嚴謹步驟：設鄰邊 $5k$、斜邊 $13k$（$k>0$），畢氏定理求對邊 $12k$。\n2. 指導隨堂 2（CASE 09）與隨堂 3（CASE 10 & 11）。", "1. 理解比值與長度區別，熟練參數 $k$ 寫法。\n2. 利用常用直角三邊比（3-4-5, 5-12-13, 7-24-25, 8-15-17）迅速計算。", "CASE 08~11\nWord 學習單\n（模組二）"),
            ("45~50 分\n(5 min)", "自檢總結\n課後任務", "1. 點擊 CASE 08 正面「自我檢核」題抽考學生。\n2. 預告下節課進行「60 秒特別角速算挑戰」，請學生先行記憶 $30^\\circ, 45^\\circ, 60^\\circ$ 幾何邊長比。", "1. 點擊卡片正面自我檢核按鈕即時檢驗。\n2. 記錄作業：完成學習單模組一與模組二前 3 題。", "正面自檢彈窗\n學習單作業")
        ]
    )

    # 4.2 第二節課
    add_lesson_section(
        2,
        "特別角推導、生活科技素養與 60 秒速算特訓 (50 分鐘)",
        "聚焦特別角幾何推導、度數弳度換算、無障礙坡道素養與速算 PK",
        [
            "幾何作圖推導 $30^\\circ, 45^\\circ, 60^\\circ$ 精確比值",
            "熟練度數與弳度量混合運算",
            "理解無障礙坡道規範並用計算機按 $\\tan^{-1}$ 反推角度",
            "在 60 秒速算破案挑戰中達成 10 題以上正確點擊"
        ],
        [
            ("00~10 分\n(10 min)", "幾何推導\n特別角速記", "1. 投影 CASE 12（正三角形對折，高為 $\\sqrt{3}$，比 $1:\\sqrt{3}:2$）與 CASE 13（等腰直角，比 $1:1:\\sqrt{2}$）。\n2. 展示 CASE 14 表 2-1：歸納 $\\sin$ 遞增與 $\\cos$ 遞減規律。\n3. 指導填寫學習單【重點 2】特別角速記表。", "1. 於草稿紙手繪兩個基本三角形，標明邊長比。\n2. 熟記規律並填妥學習單【重點 2】。", "CASE 12~14\nWord 學習單\n（重點 2）"),
            ("10~22 分\n(12 min)", "特別角運算\n弳度混合題", "1. 投影 CASE 15 & 16 精講特別角四則運算。\n2. 投影 CASE 18（弳度制運算）：強調 $\\frac{\\pi}{4}=45^\\circ, \\frac{\\pi}{6}=30^\\circ, \\frac{\\pi}{3}=60^\\circ$ 即時轉換心法。\n3. 提醒平方位置：$\\sin^2 45^\\circ = (\\sin 45^\\circ)^2 = \\frac{1}{2}$。", "1. 上台演練例題 4 與隨堂 4。\n2. 釐清 $\\sin^2\\theta$ 與 $\\sin\\theta^2$ 的根本差異。", "CASE 15~18\n白板手寫推導\nWord 學習單"),
            ("22~32 分\n(10 min)", "生活科技\n坡度與計算機", "1. 投影 CASE 19（營建署無障礙坡道高差 1 比 8 規範）：\n   · 坡度 $= \\tan\\theta = \\frac{1}{8} = 0.125$。\n   · 操作計算機：DEG 模式 $\\to$ 按 `2nd` / `SHIFT` $\\to$ `tan⁻¹` $\\to$ 輸入 0.125 $\\to$ 得 $\\theta \\approx 7.125^\\circ$。\n2. 投影 CASE 20（自行車險升坡 $20\\%$ 反推 $\\theta \\approx 11.3^\\circ$）。", "1. 拿出工程型計算機或手機計算機 App，實際依步驟操作反三角按鍵。\n2. 體會數學在公共無障礙設施與道路安全規範中的實用價值。", "CASE 19、20\n計算機操作卡\n手機 App"),
            ("32~47 分\n(15 min)", "課堂高潮\n60秒速算挑戰", "1. 投影平台頂部【阿笠博士實驗室 ‧ 60秒特別角速算挑戰】。\n2. 示範一局：題目出現 $\\sin 60^\\circ = ?$，點擊 $\\frac{\\sqrt{3}}{2}$，展示連擊 Combo。\n3. 進行全班分組 PK 賽：大螢幕與手機連線同步競速。\n4. 結算全班最高分，展示獲得之「神探徽章」。", "1. 登入手機連線頁面，全神貫注進行 60 秒限時作答。\n2. 在九宮格標準數學分數按鈕（$\\frac{1}{2}, \\frac{\\sqrt{2}}{2}, \\frac{\\sqrt{3}}{2}, \\dots$）中敏捷選擇答案。\n3. 班級氣氛高潮，徹底記熟 18 種特別角組合！", "60秒速算挑戰模組\n手機掃碼伺服器\n即時排行榜"),
            ("47~50 分\n(3 min)", "總結與特訓驗收", "1. 總結速算技巧：分母是 2 的有誰？分母是 $\\sqrt{3}$ 或 3 的是誰？\n2. 交代作業：課本 P.114 隨堂 4 與學習單生活科技題目。", "1. 口頭複誦快速回顧。\n2. 課後自主練習速算挑戰直到獲得「名偵探」級評價。", "學習單作業")
        ]
    )

    doc.add_page_break()

    # 4.3 第三節課
    add_lesson_section(
        3,
        "神級魔法六角密碼盤與四大恆等關係神技 (50 分鐘)",
        "聚焦六角密碼盤佈局、四大恆等陣法、餘角相消與齊次同除秒殺神技",
        [
            "熟記六角密碼盤「左無co、右帶co、中為1」頂點排列",
            "透過六角盤默寫對角互倒、商數、倒三角平方、橫向餘角關係",
            "靈活運用餘角轉換將非特別角化簡求值",
            "掌握「齊次式分子分母同除以 $\\cos\\theta$」秒殺一次正餘弦分式"
        ],
        [
            ("00~15 分\n(15 min)", "六角星解密\n動態探索", "1. 投影平台【神級魔法六角密碼盤】，講解佈局玄機：\n   · 上 $\\sin$，下 $\\csc$；右上 $\\cos$，右下 $\\cot$；左上 $\\tan$，左下 $\\sec$。\n   · 暗號：左邊陽光無 co、右邊陰影全帶 co、中心定海神針 1！\n2. 點擊切換三大互動模式：【對角互倒】、【相鄰商數】、【倒三角平方】。\n3. 補充橫向軸線【餘角互換】（CASE 25）。", "1. 學生於學習單【重點 3】手繪六角星盤。\n2. 學生上台點擊觸控螢幕之六角星頂點，觀察動態發光連線與解密盒方程式說明。", "魔法六角密碼盤\nCASE 21~25\nWord 學習單"),
            ("15~28 分\n(13 min)", "恆等式演練\n平方與倒數", "1. 投影 CASE 26～31 精講：\n   · 倒數求值：$\\sin 50^\\circ \\times \\csc 50^\\circ = 1$。\n   · 平方求值：$\\sin^2 40^\\circ + \\cos^2 40^\\circ = 1$。\n   · 負號陷阱：$\\sec^2 70^\\circ - \\tan^2 70^\\circ = 1$；但 $\\cot^2 12^\\circ - \\csc^2 12^\\circ = -1$！\n2. 引導以六角倒三角直觀驗算：$1+\\cot^2=\\csc^2$。", "1. 練習辨識公式移項後的正負號。\n2. 完成學習單【模組二：CASE 26~31】題目，互換批改。", "CASE 26~31\nWord 學習單\n（模組二）"),
            ("28~38 分\n(10 min)", "餘角轉化\n綜合求值", "1. 投影 CASE 32 & CASE 33（例題 7 與隨堂 7）：\n   · 觀察題幹非特別角：$40^\\circ$ 與 $50^\\circ$、$72^\\circ$ 與 $18^\\circ$。\n   · 金鑰：兩角相加為 $90^\\circ$，立即橫向換裝！\n   · $\\tan 50^\\circ \\to \\cot 40^\\circ \\implies \\tan 40^\\circ\\cot 40^\\circ = 1$。\n   · $\\csc 18^\\circ \\to \\sec 72^\\circ \\implies \\tan^2 72^\\circ - \\sec^2 72^\\circ = -1$。\n   · 原式 $= 1 + (-1) = 0$！", "1. 尋找「互餘角好夥伴」（兩角和為 $90^\\circ$）。\n2. 練習將其中一個角完全改裝為另一個角，化繁為簡。", "CASE 32、33\n白板雙色列式"),
            ("38~47 分\n(9 min)", "商數關係\n一次齊次式", "1. 投影 CASE 34 & CASE 35（已知 $\\tan\\theta=2$，求 $\\frac{\\sin\\theta+2\\cos\\theta}{3\\sin\\theta-\\cos\\theta}$）：\n   · 傳統慢速法：畫三角形設斜邊 $\\sqrt{5}$，代入通分極易算錯。\n   · 名偵探秒殺神技：分子分母同除以 $\\cos\\theta$！\n     $$\\frac{\\tan\\theta+2}{3\\tan\\theta-1} = \\frac{2+2}{3(2)-1} = \\frac{4}{5}$$\n     不需算斜邊，3 秒破案結案！", "1. 體驗「同除以 $\\cos\\theta$」秒殺法的神速與優雅。\n2. 獨立演練隨堂 8：$\\cot\\theta=\\frac{2}{3}$，同除以 $\\sin\\theta$ 求值。", "CASE 34、35\nWord 學習單\n（範例 8）"),
            ("47~50 分\n(3 min)", "總結與提煉", "1. 口訣提煉：「見乘積想倒數、見平方看倒三角、見九十度換橫向、見齊次式同除正餘弦！」\n2. 指派作業：學習單模組二範例 6~8。", "1. 齊聲複誦四大破案口訣。\n2. 整理筆記與六角星繪圖。", "學習單作業")
        ]
    )

    # 4.4 第四節課
    add_lesson_section(
        4,
        "三角恆等式深度破案、是非題辨析與劇場版考核 (50 分鐘)",
        "聚焦和差平方連環扣、是非題 5 大迷思診斷、立方和進階題與晉級名偵探授階",
        [
            "掌握 $(\\sin\\pm\\cos)^2 = 1\\pm 2\\sin\\cos$ 與 $\\tan+\\cot = \\frac{1}{\\sin\\cos}$",
            "深度診斷課本是非題 5 大迷思，能具體說出反例與錯誤根源",
            "結合因式分解立方和公式，攻克進階挑戰題（CASE 59）",
            "完成全單元學習單與少年偵探團考核，獲得名偵探勳章"
        ],
        [
            ("00~15 分\n(15 min)", "常用恆等式\n和差連環扣", "1. 投影 CASE 36 推導：\n   · $(\\sin\\pm\\cos)^2 = 1\\pm 2\\sin\\theta\\cos\\theta$。\n   · $\\tan+\\cot = \\frac{\\sin}{\\cos}+\\frac{\\cos}{\\sin} = \\frac{1}{\\sin\\theta\\cos\\theta}$。\n2. 投影 CASE 37 & 38（例題 9：已知 $\\sin+\\cos=\\frac{4}{3}$）：\n   · 兩邊平方求得 $\\sin\\cos = \\frac{7}{18}$。\n   · 取倒數立即破案 $\\tan+\\cot = \\frac{18}{7}$！\n3. 指導隨堂 9（CASE 39 & 40：差平方）。", "1. 筆記口訣：「見和差必平方、見 $\\tan+\\cot$ 必通分倒數！」。\n2. 練習和平方展開與求倒數運算。", "CASE 36~40\nWord 學習單\n（模組二）"),
            ("15~28 分\n(13 min)", "是非題專題\n五大迷思診斷", "1. 投影 CASE 41～45（課本 P.123 觀念是非題五大案）：\n   · 案 1 (○)：已知銳角 $\\sin\\theta$，可求其餘五者（畢氏定理定三邊比唯一確定）。\n   · 案 2 (✕)：$\\sin A=\\frac{3}{5}$ 能否求三邊長？（比值非長度，無數據不可求）。\n   · 案 3 (✕)：$\\sin^2\\alpha+\\cos^2\\beta=1$？（角不同不成立，平方限定同角）。\n   · 案 4 (✕)：$\\sin\\theta\\times\\cos\\theta=1$？（倒數應為 $\\sin\\csc=1$）。\n   · 案 5 (✕)：$(\\sin+\\cos)^2=1$？（遺漏交叉項 $2\\sin\\cos$）。\n2. 引導填寫學習單【模組三：是非題專題】。", "1. 於座位上以手勢比出「○」或「✕」。\n2. 每題指名一位學生說出「錯誤原因與反例」，徹底掃除概念盲區。", "CASE 41~45 是非圖解卡\n正誤紅印章\nWord 學習單"),
            ("28~40 分\n(12 min)", "進階挑戰\n立方和因式分解", "1. 投影 CASE 56（進階 8：通分平方化簡求值）。\n2. 投影 CASE 57～59（進階 9 系列 ‧ 劇場版終極破案）：\n   · 已知 $\\tan x+\\cot x = \\frac{25}{12}$。\n   · (1) $\\sin x\\cos x = \\frac{12}{25}$ ； (2) $\\sin x+\\cos x = \\frac{7}{5}$。\n   · (3) 求 $\\sin^3 x + \\cos^3 x$？\n   · 立方和公式 $a^3+b^3=(a+b)(1-ab)$！\n   · 代入 $(\\frac{7}{5})(1-\\frac{12}{25}) = \\frac{7}{5}\\times\\frac{13}{25} = \\frac{91}{125}$！", "1. 結合因式分解立方和公式與高中三角比，挑戰高階思考。\n2. 體驗一題連環三小題、層層抽絲剝繭破案的成就感！", "CASE 56~59\nWord 學習單\n（模組四）"),
            ("40~47 分\n(7 min)", "總結神技\n偵探團授階", "1. 投影 CASE 60【名偵探破案密碼 ‧ 終極神技卷】：六大心法總驗收。\n2. 開啟【少年偵探團升級考核】模組，引導完成 STAGE 01～06 總檢核。\n3. 頒發「技術高中名偵探」榮譽勳章。", "1. 總覽全單元 60 案架構。\n2. 於手機或大螢幕完成闖關總測驗，檢視個人總積分與排名。", "CASE 60 神技卷\n闖關總模組\n結業勳章"),
            ("47~50 分\n(3 min)", "收卷與反思", "1. 收回完成之 JIS B4 雙欄學習單。\n2. 提醒學生將網站捷徑存至手機主畫面複習。", "1. 填寫學習單最後一欄【偵探搜查反思筆記】。\n2. 繳交學習單。", "學習單收件")
        ]
    )

    doc.add_page_break()

    # 5. 參、多元評量策略與評分規準 (Rubrics)
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    hr3 = h3.add_run("參、多元評量策略與學習單評分規準 (Rubrics)")
    hr3.font.name = "微軟正黑體"
    hr3.font.size = Pt(13)
    hr3.bold = True
    hr3.font.color.rgb = RGBColor(30, 58, 138)

    rub_intro = doc.add_paragraph()
    rub_intro.paragraph_format.space_after = Pt(6)
    add_math_runs(rub_intro, "本單元落實「歷程評量」與「表現任務評量」，配分比重：課堂數位速算互動（30%）、JIS B4 雙欄學習單書寫（40%）、少年偵探團 STAGE 闖關總評量（30%）。", default_color=(51, 65, 85), font_size=Pt(9.5))

    rub_tbl = doc.add_table(rows=5, cols=5)
    rub_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(rub_tbl, "CBD5E1")
    rub_widths = [Mm(26), Mm(36), Mm(36), Mm(38), Mm(38)]
    rub_hdrs = ["評分維度", "A 優異 (90-100分)", "B 良好 (80-89分)", "C 基礎 (70-79分)", "D 待加強 (<70分)"]
    
    for c_idx, h_text in enumerate(rub_hdrs):
        cell = rub_tbl.cell(0, c_idx)
        cell.width = rub_widths[c_idx]
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        set_cell_shading(cell, "1E3A8A")
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(0)
        r = cp.add_run(h_text)
        r.font.name = "微軟正黑體"
        r.font.size = Pt(9)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    rub_data = [
        ("定義掌握度", "六大三角函數定義清晰，能靈活切換視角，設參數 $k$ 規範嚴謹。", "正確認出六大三角比，視角切換偶有遲疑但能自我更正。", "能認出前三大三角比，伴隨函數（餘割、正割、餘切）常混淆。", "無法正確認出對邊與鄰邊，三角比定義混亂。"),
        ("特別角運算", "特別角度數與弳度值反射性寫出，四則運算精準無誤。", "能推導特別角數值，混合運算過程偶有通分計算失誤。", "需仰賴繪製幾何圖形方能寫出數值，弳度量不熟練。", "特別角數值記錯或完全無法作答。"),
        ("恆等式靈活度", "熟練默寫六角盤四大關係，秒殺齊次分式與和差連環扣。", "能使用基本平方與倒數關係，齊次式同除經提示後能完成。", "僅能使用 $\\sin^2+\\cos^2=1$，無法掌握餘角轉換與齊次式神技。", "恆等式公式嚴重混淆，無法進行代數化簡。"),
        ("是非題診斷力", "5 題是非題全對，且能以嚴謹數學反例具體說明錯誤原因。", "是非題答對 4 題，能說明主要錯誤概念。", "是非題答對 3 題，僅知對錯但說不出具體原因。", "缺乏概念辨析能力，猜題成分居多。")
    ]
    
    for r_idx, r_vals in enumerate(rub_data, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, text in enumerate(r_vals):
            cell = rub_tbl.cell(r_idx, c_idx)
            cell.width = rub_widths[c_idx]
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_shading(cell, bg)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            if c_idx == 0:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_math_runs(cp, text, default_color=(120, 53, 15), bold=True, font_size=Pt(8.5))
            else:
                add_math_runs(cp, text, default_color=(51, 65, 85), font_size=Pt(8.5))
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. 肆、差異化教學與班級經營指引
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    hr4 = h4.add_run("肆、差異化教學、情境融入與班級經營指引")
    hr4.font.name = "微軟正黑體"
    hr4.font.size = Pt(13)
    hr4.bold = True
    hr4.font.color.rgb = RGBColor(30, 58, 138)

    tips = [
        ("1. 拔尖（學習優異組）", "引導攻克 CASE 50（雙括號弳度多項式相乘）、CASE 52（角度平方消去法）與 CASE 59（立方和因式分解代入）。邀請擔任「名偵探小隊長」，於分組活動中協助同儕診斷是非題錯誤原因。"),
        ("2. 扶弱（學習緩步組）", "專注掌握 CASE 01～07（定義與邊長比）與 CASE 12～14（特別角三大金剛）。鼓勵反覆操作「阿笠博士測距儀」與「60秒速算挑戰」，透過視覺幾何與遊戲化反饋降低數學焦慮；是非題重點輔導「比值非長度」與「同角平方才為 1」。"),
        ("3. 探案沉浸感塑造", "上課時稱呼學生為「特派搜查員」，將解題稱為「蒐集線索、緝拿真兇」。每完成一個例題，點擊卡片正面「自我檢核」，答對時全班齊聲喊：「CASE CLOSED！真相永遠只有一個！」，激發團隊榮譽感。"),
        ("4. 投影白板互動", "使用全螢幕彈窗劇院頂部之「偵探筆記」白板工具，用領結紅筆標記對邊、西裝藍筆標記鄰邊、眼鏡金筆標記斜邊，顏色一致性加深空間幾何記憶。")
    ]
    
    for t_title, t_desc in tips:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(2)
        tp.paragraph_format.space_after = Pt(2)
        add_math_runs(tp, t_title + "：", default_color=(180, 83, 9), bold=True, font_size=Pt(9.5))
        add_math_runs(tp, t_desc, default_color=(51, 65, 85), font_size=Pt(9.5))

    out_file = "2-2_三角函數的定義與性質_四節課教學計畫與教案.docx"
    doc.save(out_file)
    print(f"成功產出正式教案 Word 文件: {out_file} (檔案大小: {os.path.getsize(out_file):,} 位元組)")
    return out_file

if __name__ == "__main__":
    make_lesson_plan_doc()
