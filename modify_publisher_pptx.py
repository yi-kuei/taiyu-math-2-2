# -*- coding: utf-8 -*-
"""
泰宇高中數學 2-2「三角函數的定義與性質」原廠簡報全方位修訂增強腳本
精準定位版：按照原始 77 頁投影片精準修改、插入先備知識與導航頁，並注入完整教學引導與備忘錄。
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC_FILE = "技高數學C（一）-CH2(2-2)_8201B(41)_原始檔備份.pptx"
OUT_FILE = "技高數學C（一）-CH2(2-2)_8201B(41).pptx"
COPY_FILE = "2-2_三角函數的定義與性質_原廠修訂教學版.pptx"

# 品牌教學色系
NAVY_BLUE = RGBColor(15, 34, 64)          # 主標題深海軍藍 #0F2240
TEAL = RGBColor(14, 116, 144)             # 亮青藍 #0E7490
AMBER_GOLD = RGBColor(180, 83, 9)         # 琥珀金 #B45309
BG_LIGHT_BLUE = RGBColor(241, 245, 249)   # 淺藍灰底色
BG_LIGHT_YELLOW = RGBColor(254, 252, 232) # 溫暖淺黃底色
DARK_TEXT = RGBColor(30, 41, 59)          # 內文深灰 #1E293B
CRIMSON = RGBColor(185, 28, 28)           # 強調深紅 #B91C1C
WHITE = RGBColor(255, 255, 255)
BORDER_GRAY = RGBColor(203, 213, 225)     # 邊框灰
BORDER_GOLD = RGBColor(245, 158, 11)      # 琥珀金邊框
BORDER_BLUE = RGBColor(59, 130, 246)      # 藍色邊框

def format_run(run, font_name="微軟正黑體", size_pt=18, bold=False, color=DARK_TEXT):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color

def add_header(slide, title_text, subtitle_text=""):
    """在投影片頂部添加統一大標題"""
    top_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.9))
    tf = top_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.LEFT
    format_run(p.runs[0], font_name="微軟正黑體", size_pt=26, bold=True, color=NAVY_BLUE)
    
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.alignment = PP_ALIGN.LEFT
        format_run(p2.runs[0], font_name="微軟正黑體", size_pt=14, bold=False, color=TEAL)

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_GRAY):
    """建立一張圓角矩形卡片容器"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def set_speaker_notes(slide, notes_text):
    """設置講者備忘錄"""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text.strip()

def add_guidance_box(slide, left, top, width, height, title, points, bg_color=BG_LIGHT_YELLOW, border_color=BORDER_GOLD):
    """在題目投影片中添加解題重要方向引導框"""
    card = add_card(slide, left, top, width, height, bg_color=bg_color, border_color=border_color)
    tb = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.36), height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    # 標題
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.alignment = PP_ALIGN.LEFT
    format_run(p0.runs[0], font_name="微軟正黑體", size_pt=16, bold=True, color=CRIMSON)
    
    # 重點條列
    for pt_text in points:
        p = tf.add_paragraph()
        p.text = pt_text
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3)
        format_run(p.runs[0], font_name="微軟正黑體", size_pt=13, bold=False, color=DARK_TEXT)

def create_prerequisite_slide_1(slide):
    add_header(slide, "【先備知識 ①】直角三角形與畢氏定理（奠定邊長計算基石）", "溫故知新 ‧ 回顧國中幾何基礎，無縫接軌高職三角函數")
    
    c1 = add_card(slide, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.7), bg_color=WHITE, border_color=BORDER_BLUE)
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "📐 畢氏定理（Pythagorean Theorem）"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    p = tf1.add_paragraph()
    p.text = "在直角三角形 ABC 中，若 ∠C = 90°，則："
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=15, color=DARK_TEXT)
    
    p = tf1.add_paragraph()
    p.text = "   a² + b² = c²  (兩股平方和 ＝ 斜邊平方)"
    p.space_before = Pt(8)
    format_run(p.runs[0], size_pt=18, bold=True, color=CRIMSON)
    
    p = tf1.add_paragraph()
    p.text = "💡 快速運算心法："
    p.space_before = Pt(12)
    format_run(p.runs[0], size_pt=16, bold=True, color=AMBER_GOLD)
    
    p = tf1.add_paragraph()
    p.text = "• 求斜邊 c ＝ √(a² + b²) （兩股平方相加開根號）"
    p.space_before = Pt(4)
    format_run(p.runs[0], size_pt=14, color=DARK_TEXT)
    
    p = tf1.add_paragraph()
    p.text = "• 求對邊或鄰邊 ＝ √(斜邊² - 另邊²) （斜邊平方減另邊平方）"
    p.space_before = Pt(4)
    format_run(p.runs[0], size_pt=14, color=DARK_TEXT)
    
    p = tf1.add_paragraph()
    p.text = "⚡ 課堂暖身提問："
    p.space_before = Pt(12)
    format_run(p.runs[0], size_pt=15, bold=True, color=TEAL)
    
    p = tf1.add_paragraph()
    p.text = "「直角三角形斜邊長為 13，其中一股長為 5，另一股長秒答是多少？」"
    p.space_before = Pt(4)
    format_run(p.runs[0], size_pt=14, bold=True, color=DARK_TEXT)

    c2 = add_card(slide, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.7), bg_color=BG_LIGHT_YELLOW, border_color=BORDER_GOLD)
    tb2 = slide.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "⭐ 技高統測必背：常用整數邊長比"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    p = tf2.add_paragraph()
    p.text = "1. ( 3 ,  4 ,  5 )  及其倍數 ( 6, 8, 10 )、( 9, 12, 15 )\n2. ( 5 , 12 , 13 )  ★ 統測最愛考！\n3. ( 7 , 24 , 25 )\n4. ( 8 , 15 , 17 )"
    p.space_before = Pt(8)
    format_run(p.runs[0], size_pt=15, color=DARK_TEXT)
    
    p = tf2.add_paragraph()
    p.text = "🎯 兩大黃金特殊角邊長比例："
    p.space_before = Pt(14)
    format_run(p.runs[0], size_pt=17, bold=True, color=CRIMSON)
    
    p = tf2.add_paragraph()
    p.text = "• 45° - 45° - 90° 直角三角形：\n   邊長比 ＝ 1 : 1 : √2 （等腰直角三角形）"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=14, color=DARK_TEXT)
    
    p = tf2.add_paragraph()
    p.text = "• 30° - 60° - 90° 直角三角形：\n   邊長比 ＝ 1 : √3 : 2 （30°對邊最短短股為1，斜邊為2）"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=14, color=DARK_TEXT)
    
    set_speaker_notes(slide, """【教學焦點】：
1. 引起動機：三角函數本質就是「直角三角形邊長比」，因此對畢氏定理與邊長比的熟練度直接決定解題速度。
2. 課堂互動：點名學生抽問「3-4-5」、「5-12-13」，並提醒 30°-60°-90° 的 √3 是在 60° 的對邊（即 30° 的鄰邊），千萬不要把斜邊和鄰邊搞混。
3. 提示：考試時遇到分數或根號時，不要慌張，隨時畫出標準直角三角形草圖進行邊長標註。""")

def create_prerequisite_slide_2(slide):
    add_header(slide, "【先備知識 ②】相似形與三角比不變性（三角函數的起源）", "核心探究 ‧ 為什麼直角三角形放大縮小，邊長比卻完全不變？")
    
    c1 = add_card(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.3), bg_color=WHITE, border_color=BORDER_BLUE)
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.3), Inches(2.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "🔍 AA 相似定理：幾何圖形縮放不影響邊長比值"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    p = tf1.add_paragraph()
    p.text = "在任意兩個直角三角形中，只要它們具有相同的銳角 θ（例如 ∠A = θ，∠C = 90°），\n由國中 AA 相似定理可知，這兩個直角三角形「形狀完全相同（相似）」，僅大小不同。"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=15, color=DARK_TEXT)
    
    p = tf1.add_paragraph()
    p.text = "👉 結論：無論三角形被放大 2 倍、10 倍或縮小，其「對應邊長的比值」永遠是固定定值！"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=16, bold=True, color=CRIMSON)

    c2 = add_card(slide, Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.2), bg_color=BG_LIGHT_BLUE, border_color=TEAL)
    tb2 = slide.shapes.add_textbox(Inches(1.0), Inches(4.15), Inches(11.3), Inches(1.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "✨ 角度決定比例 ‧ 三角函數的誕生"
    format_run(p.runs[0], size_pt=18, bold=True, color=TEAL)
    
    p = tf2.add_paragraph()
    p.text = "既然比值不受三角形大小影響，只由「銳角 θ 的大小」唯一決定，數學家就為這些比值取了名字：\n• (對邊 / 斜邊) ＝ 正弦 (sine, sin θ)\n• (鄰邊 / 斜邊) ＝ 餘弦 (cosine, cos θ)\n• (對邊 / 鄰邊) ＝ 正切 (tangent, tan θ)"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=15, color=DARK_TEXT)
    
    p = tf2.add_paragraph()
    p.text = "💡 「三角函數」是連接【角度】與【邊長長度】的數學橋樑，廣泛應用於建築測量、航海與工程學！"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=15, bold=True, color=AMBER_GOLD)

    set_speaker_notes(slide, """【教學焦點】：
1. 破除迷思：許多學生誤以為直角三角形變大，sin 值就會變大。務必在此頁強調：相似三角形對應邊成比例，分子分母同乘 k 倍後 k 會消掉，比值只跟角度 θ 有關！
2. 啟發思考：問學生「如果在操場畫一個 30° 直角三角形，和在課本畫一個 30° 直角三角形，對邊/斜邊的比值一樣嗎？」引導全班齊聲回答「一樣，都是 1/2」！""")

def create_roadmap_slide(slide):
    add_header(slide, "【單元核心學習地圖】2-2 四大核心模組全景導覽", "掌握全章學習脈絡 ‧ 清楚知道每一堂課要學什麼")
    
    modules = [
        ("模組 1：銳角的三角函數 (2-2.1)", 
         "• 直角三角形中 6 個邊長比定義 (sin, cos, tan, cot, sec, csc)\n• 英文草寫筆順記憶法與正餘對應\n• 「一角求六角」：已知一角三角比，求其餘五者"),
        
        ("模組 2：特別角三角函數值 (2-2.2)", 
         "• 30°、45°、60° 特別角精確函數值\n• 特殊角度邊長比幾何速推\n• 特別角三角函數值的代數四則混和運算"),
        
        ("模組 3：三角函數恆等關係 (2-2.3)", 
         "• 倒數關係：sin·csc=1, cos·sec=1, tan·cot=1\n• 商數關係：tan=sin/cos, cot=cos/sin\n• 平方關係：sin²+cos²=1, 1+tan²=sec², 1+cot²=csc²"),
        
        ("模組 4：三角恆等式化簡與證明 (2-2.4)", 
         "• 齊次式同除以 cosθ 技巧（秒殺統測題）\n• 「化為正餘弦」統一變數法\n• 乘法公式展開：(sinθ ± cosθ)² ＝ 1 ± 2sinθcosθ")
    ]
    
    positions = [
        (Inches(0.8), Inches(1.5), Inches(5.6), Inches(2.2)),
        (Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.2)),
        (Inches(0.8), Inches(3.9), Inches(5.6), Inches(2.3)),
        (Inches(6.8), Inches(3.9), Inches(5.7), Inches(2.3))
    ]
    
    for i, (title, content) in enumerate(modules):
        left, top, w, h = positions[i]
        c = add_card(slide, left, top, w, h, bg_color=WHITE, border_color=TEAL if i < 2 else AMBER_GOLD)
        tb = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), w - Inches(0.36), h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        format_run(p.runs[0], size_pt=16, bold=True, color=NAVY_BLUE)
        
        for line in content.split('\n'):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.space_before = Pt(3)
            format_run(p_line.runs[0], size_pt=13, color=DARK_TEXT)

    set_speaker_notes(slide, """【教學焦點】：
1. 學習地圖總覽：向學生展示 2-2 單元的完整學習脈絡。這四大模組層層遞進：
   - 模組 1 給出工具定義（6 個函數）；
   - 模組 2 給出具體數值（30°, 45°, 60°）；
   - 模組 3 找出內在結構（恆等關係）；
   - 模組 4 進入統測實戰（代數化簡與證明）。
2. 給學生的信心喊話：只要跟著每堂課的節奏，三角函數是統測最好拿分的單元！""")

def create_memory_guide_slide(slide):
    add_header(slide, "【核心記憶法】六大三角函數英文字母筆順記憶圖鑑", "看圖秒記 ‧ 告別死記硬背，英文字母筆順帶你直達邊長比")
    
    c1 = add_card(slide, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.7), bg_color=WHITE, border_color=BORDER_BLUE)
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "✍️ 三大基礎函數（筆順記憶法）"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    funcs = [
        ("sin θ ＝ 對邊 / 斜邊", "草寫『s』筆順：從斜邊出發，劃向對邊。\n（口訣：正弦對斜邊，s 劃對邊）"),
        ("cos θ ＝ 鄰邊 / 斜邊", "草寫『c』筆順：從斜邊出發，包夾角 θ 到鄰邊。\n（口訣：餘弦包夾角，c 夾斜鄰邊）"),
        ("tan θ ＝ 對邊 / 鄰邊", "草寫『t』筆順：先走鄰邊橫線，再直劃對邊。\n（口訣：正切直直立，t 由鄰到對）")
    ]
    
    for f_title, f_desc in funcs:
        p_t = tf1.add_paragraph()
        p_t.text = f"• {f_title}"
        p_t.space_before = Pt(8)
        format_run(p_t.runs[0], size_pt=16, bold=True, color=CRIMSON)
        
        p_d = tf1.add_paragraph()
        p_d.text = f_desc
        p_d.space_before = Pt(2)
        format_run(p_d.runs[0], size_pt=13, color=DARK_TEXT)

    c2 = add_card(slide, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.7), bg_color=BG_LIGHT_YELLOW, border_color=BORDER_GOLD)
    tb2 = slide.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "🔄 三大倒數函數（對稱記憶法）"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    recips = [
        ("csc θ ＝ 斜邊 / 對邊 ＝ 1 / sin θ", "餘割 (cosecant)：正弦 sin 的倒數。\n（注意：有 co 的 csc 是無 co 的 sin 倒數）"),
        ("sec θ ＝ 斜邊 / 鄰邊 ＝ 1 / cos θ", "正割 (secant)：餘弦 cos 的倒數。\n（注意：無 co 的 sec 是有 co 的 cos 倒數）"),
        ("cot θ ＝ 鄰邊 / 對邊 ＝ 1 / tan θ", "餘切 (cotangent)：正切 tan 的倒數。\n（正切餘切名字對應，互為倒數最直觀）")
    ]
    
    for r_title, r_desc in recips:
        p_t = tf2.add_paragraph()
        p_t.text = f"• {r_title}"
        p_t.space_before = Pt(8)
        format_run(p_t.runs[0], size_pt=15, bold=True, color=TEAL)
        
        p_d = tf2.add_paragraph()
        p_d.text = r_desc
        p_d.space_before = Pt(2)
        format_run(p_d.runs[0], size_pt=13, color=DARK_TEXT)
        
    p_sum = tf2.add_paragraph()
    p_sum.text = "📢 口訣：『正餘互配無 co 對帶 co，正切餘切本一家！』"
    p_sum.space_before = Pt(10)
    format_run(p_sum.runs[0], size_pt=14, bold=True, color=AMBER_GOLD)

    set_speaker_notes(slide, """【教學焦點】：
1. 肢體與筆順教學：在黑板上帶學生用手在空中劃出草寫 s、c、t 的筆順，感受起筆邊（分母）到落筆邊（分子）的流向。
2. 易錯提醒：學生最容易把 csc 和 sec 的倒數對象搞反（誤以為 csc 配 cos）。在此特別強調「一個有 co、一個沒 co」的反向配對原則！""")

def create_super_hexagon_slide(slide):
    add_header(slide, "【恆等關係核心神器】三角函數超級六邊形（Super Hexagon）", "一圖破全章 ‧ 倒數、商數、平方三大恆等關係視覺化速記神圖")
    
    c1 = add_card(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(4.7), bg_color=WHITE, border_color=BORDER_BLUE)
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.4), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "⚡ 超級六邊形三大神奇法則"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    rules = [
        ("① 倒數關係（穿心對角線）", "穿過中心點 1 的對角線兩端頂點相乘等於 1：\n• sin θ · csc θ ＝ 1\n• cos θ · sec θ ＝ 1\n• tan θ · cot θ ＝ 1"),
        ("② 商數關係（相鄰頂點比）", "任一頂點等於順時針或逆時針相鄰兩頂點之商：\n• tan θ ＝ sin θ / cos θ\n• cot θ ＝ cos θ / sin θ"),
        ("③ 平方關係（倒立正三角形）", "三個倒立三角形中，上方兩頂點平方和等於下方頂點：\n• sin² θ + cos² θ ＝ 1\n• 1 + tan² θ ＝ sec² θ\n• 1 + cot² θ ＝ csc² θ")
    ]
    
    for r_t, r_c in rules:
        p_t = tf1.add_paragraph()
        p_t.text = r_t
        p_t.space_before = Pt(8)
        format_run(p_t.runs[0], size_pt=16, bold=True, color=CRIMSON)
        
        p_c = tf1.add_paragraph()
        p_c.text = r_c
        p_c.space_before = Pt(2)
        format_run(p_c.runs[0], size_pt=13, color=DARK_TEXT)

    c2 = add_card(slide, Inches(7.0), Inches(1.5), Inches(5.5), Inches(4.7), bg_color=BG_LIGHT_YELLOW, border_color=BORDER_GOLD)
    tb2 = slide.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.1), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "🎯 六邊形頂點配置心法"
    format_run(p.runs[0], size_pt=18, bold=True, color=NAVY_BLUE)
    
    p = tf2.add_paragraph()
    p.text = "• 右半部：全部都是帶『co』的餘函數！\n   (右上 cos θ、正右 cot θ、右下 csc θ)"
    p.space_before = Pt(8)
    format_run(p.runs[0], size_pt=14, bold=True, color=CRIMSON)
    
    p = tf2.add_paragraph()
    p.text = "• 左半部：全部都是無『co』的正函數！\n   (左上 sin θ、正左 tan θ、左下 sec θ)"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=14, bold=True, color=TEAL)
    
    p = tf2.add_paragraph()
    p.text = "• 正中心：寫上醒目的數字 『 1 』！"
    p.space_before = Pt(6)
    format_run(p.runs[0], size_pt=14, bold=True, color=DARK_TEXT)
    
    p = tf2.add_paragraph()
    p.text = "💡 課堂實戰應用技巧："
    p.space_before = Pt(12)
    format_run(p.runs[0], size_pt=16, bold=True, color=AMBER_GOLD)
    
    p = tf2.add_paragraph()
    p.text = "遇到三角恆等式求值或化簡時，在草稿紙邊緣花 10 秒鐘畫出這個六邊形，所有 8 個恆等式一目了然，絕不混淆忘記！"
    p.space_before = Pt(4)
    format_run(p.runs[0], size_pt=13, color=DARK_TEXT)

    set_speaker_notes(slide, """【教學焦點】：
1. 六邊形法寶：此圖為全世界數學教育界推崇之「三角函數終極記憶法」。
2. 課堂帶領：請同學拿出筆在筆記本上跟著畫一個正六邊形，標出中央的 1，以及左三無 co、右三有 co。
3. 驗證平方關係：特別指出三個倒三角形：
   - 頂部倒三角：sin² + cos² = 1
   - 左下倒三角：tan² + 1 = sec²
   - 右下倒三角：1 + cot² = csc²
4. 強調統測必備，化簡題目直接秒殺！""")

def run_modifications():
    print("Loading original presentation...")
    prs = Presentation(SRC_FILE)
    
    # 1. 修正 Slide 2（章節封面）錯字
    s2 = prs.slides[1]
    for sh in s2.shapes:
        if sh.has_text_frame and "2-2有向角及其度量" in sh.text_frame.text:
            sh.text_frame.text = "2-2 三角函數的定義與性質"
            p = sh.text_frame.paragraphs[0]
            format_run(p.runs[0], size_pt=40, bold=True, color=RGBColor(255, 255, 255))
            print("Fixed Slide 2 title successfully!")
    
    set_speaker_notes(s2, """【本課導言】：
本章節為技高數學三角學最核心單元『2-2 三角函數的定義與性質』。學生在國中已學過畢氏定理與特殊直角三角形邊長比，本單元將把幾何邊長比提升為以『銳角』為自變數的六個三角函數，並建立特別角值與三大恆等關係。""")

    # 2. 在原廠投影片進行例題、隨堂練習、習題的教學引導與備忘錄增強（直接在原始 77 頁中精確處理）
    # 原始索引對照表：
    # Slide 15 (idx 14): 例題 1
    # Slide 16 (idx 15): 隨堂練習 1
    # Slide 18 (idx 17): 例題 2
    # Slide 20 (idx 19): 隨堂練習 2
    # Slide 21 (idx 20): 例題 3
    # Slide 22 (idx 21): 隨堂練習 3
    # Slide 31 (idx 30): 例題 4
    # Slide 33 (idx 32): 隨堂練習 4
    # Slide 35 (idx 34): 例題 5
    # Slide 41 (idx 40): 隨堂練習 5
    # Slide 55 (idx 54): 例題 6
    # Slide 56 (idx 55): 隨堂練習 6
    # Slide 57 (idx 56): 例題 7
    # Slide 58 (idx 57): 隨堂練習 7
    # Slide 59 (idx 58): 例題 8
    # Slide 60 (idx 59): 隨堂練習 8
    # Slide 65 (idx 64): 例題 9
    # Slide 66 (idx 65): 隨堂練習 9
    # Slide 68~77 (idx 67~76): 習題 1~10

    teaching_data = {
        14: { # 例題 1
            "title": "🎯 例題 1 課堂解題思考方向（求六個三角函數值）：",
            "points": [
                "① 觀察已知邊：確認兩股為 3、4，利用畢氏定理求斜邊 c ＝ √(3²+4²) ＝ 5",
                "② 標明基準角：以 ∠A 為基準，對邊＝3、鄰邊＝4、斜邊＝5",
                "③ 套用定義：先求 sin A＝3/5、cos A＝4/5、tan A＝3/4",
                "④ 倒數對應：立即寫出 csc A＝5/3、sec A＝5/4、cot A＝4/3"
            ],
            "notes": """【課堂解題引導】：
1. 提問：這題要算六個三角函數，第一步必須先知道什麼？（引導回答：直角三角形的三邊長！）
2. 畢氏定理速算：斜邊 ＝ √(3² + 4²) ＝ 5。
3. 基準角對照：提醒同學視線站在 ∠A，對面是 3（對邊），腳底下是 4（鄰邊），最長是 5（斜邊）。
4. 倒數檢驗：算完前三個，後三個直接取倒數，大幅節省時間。"""
        },
        15: { # 隨堂練習 1
            "title": "🎯 隨堂練習 1 課堂思考方向（基準角切換）：",
            "points": [
                "① 基準角切換：注意本題改求 ∠B 的三角函數值！",
                "② 邊長角色互換：∠B 的對邊變為 4、鄰邊變為 3、斜邊仍為 5",
                "③ 餘角關係觀察：sin B ＝ cos A、cos B ＝ sin A、tan B ＝ cot A"
            ],
            "notes": """【課堂解題引導】：
1. 核心盲點：學生最常犯的錯誤是直接抄剛才 ∠A 的答案！
2. 強調觀念：當基準角從 ∠A 換到 ∠B，對邊與鄰邊的角色立刻對調。
3. 引導學生發現：sin B 剛好等於 cos A，為後續餘角關係埋下伏筆。"""
        },
        17: { # 例題 2
            "title": "🎯 例題 2 課堂思考方向（一角求六角經典題）：",
            "points": [
                "① 設比例邊長：已知 sin A ＝ 5/13，令對邊為 5、斜邊為 13",
                "② 畢氏定理求第三邊：鄰邊 ＝ √(13² - 5²) ＝ 12",
                "③ 繪製輔助直角三角形：標上邊長 5, 12, 13",
                "④ 依定義直接寫出其餘五個三角函數值"
            ],
            "notes": """【課堂解題引導】：
1. 破題關鍵：已知一個三角函數值，就能立刻還原出一個直角三角形的兩邊比例！
2. 畫圖習慣：要求每位學生在考卷邊緣隨手畫一個小直角三角形，標上 5 與 13。
3. 鄰邊計算：13² - 5² = 169 - 25 = 144 = 12²，秒出 12。
4. 叮嚀：題目問其餘各函數值，不包含 sin A 本身。"""
        },
        19: { # 隨堂練習 2
            "title": "🎯 隨堂練習 2 課堂思考方向：",
            "points": [
                "① 已知 tan A ＝ 3/4，即對邊與鄰邊之比為 3 : 4",
                "② 畢氏定理求斜邊：斜邊 ＝ √(3² + 4²) ＝ 5",
                "③ 依序寫出其餘五個三角函數值"
            ],
            "notes": """【課堂解題引導】：
1. 提問：tan 是什麼邊比什麼邊？（對邊/鄰邊）。
2. 連續計算：對邊 3、鄰邊 4，斜邊必為 5。其餘函數直接秒殺。"""
        },
        20: { # 例題 3
            "title": "🎯 例題 3 課堂思考方向（多直角三角形邊長求解）：",
            "points": [
                "① 尋找關鍵共用邊：連接兩個直角三角形的橋樑是哪一個線段？",
                "② 第一步求解：在已知條件較多的三角形中，利用三角比求出共用邊",
                "③ 第二步推進：將共用邊代入第二個直角三角形，求解未知目標邊長"
            ],
            "notes": """【課堂解題引導】：
1. 幾何破題：引導學生看圖，找出兩個直角三角形的交界邊。
2. 模組化計算：先封閉看第一個三角形，求出中間變數；再切換到第二個三角形求最終答案。"""
        },
        21: { # 隨堂練習 3
            "title": "🎯 隨堂練習 3 課堂思考方向：",
            "points": [
                "① 依圖形標示已知邊長與特殊角",
                "② 利用共用邊進行線段長度換算",
                "③ 求出未知目標線段長度"
            ],
            "notes": """【課堂解題引導】：
引導學生先看大三角形與小三角形的幾何關聯，找出突破口。"""
        },
        30: { # 例題 4
            "title": "🎯 例題 4 課堂思考方向（特別角混和四則運算）：",
            "points": [
                "① 熟練查表/手勢推導：準確代入 30°、45°、60° 函數值",
                "② 運算順序叮嚀：先算『三角函數的平方』，再算乘除，最後算加減",
                "③ 注意：sin² 30° 代表 (sin 30°)² ＝ (1/2)² ＝ 1/4"
            ],
            "notes": """【課堂解題引導】：
1. 易錯提醒：sin² 30° 不是 sin(900°)！一定要向學生強調指數位置的記法。
2. 逐步代入：請同學在黑板上一一寫出各特別角數值，再進行分數通分加減。"""
        },
        32: { # 隨堂練習 4
            "title": "🎯 隨堂練習 4 課堂思考方向：",
            "points": [
                "① 查表帶入各特殊角之數值",
                "② 優先計算各項平方值",
                "③ 依序進行通分加減與乘除運算"
            ],
            "notes": """【課堂解題引導】：
讓學生上台板書，確認其特別角數值記憶是否準確無誤。"""
        },
        34: { # 例題 5
            "title": "🎯 例題 5 課堂思考方向（工程坡道度數應用）：",
            "points": [
                "① 坡度定義：坡度 ＝ 垂直上升高度 / 水平距離 ＝ tan θ",
                "② 依題意：tan θ ＝ 1/8 ＝ 0.125",
                "③ 反三角函數查表/計算機：求出對應夾角 θ 約為多少度"
            ],
            "notes": """【課堂解題引導】：
1. 生活素養：將坡度與正切函數 tanθ 連結，說明工程設計之規範依據。
2. 介紹計算機 tan⁻¹ 反函數的操作方法。"""
        },
        40: { # 隨堂練習 5
            "title": "🎯 隨堂練習 5 課堂思考方向：",
            "points": [
                "① 險升坡百分比：20% ＝ 每前進 100 公尺上升 20 公尺",
                "② 坡度正切值：tan θ ＝ 20 / 100 ＝ 0.2",
                "③ 反推夾角：求對應夾角 θ 大約為多少度"
            ],
            "notes": """【課堂解題引導】：
生活情境連結：引導學生計算自行車公路常見的坡度標誌意義。"""
        },
        54: { # 例題 6
            "title": "🎯 例題 6 課堂思考方向（倒數與商數關係化簡）：",
            "points": [
                "① 統一變數原則：遇到切割函數（tan, cot, sec, csc），全化為 sin 與 cos",
                "② 商數關係代換：tan θ ＝ sin θ / cos θ，cot θ ＝ cos θ / sin θ",
                "③ 通分合併式子：觀察分子分母是否能因式分解或公因式約分"
            ],
            "notes": """【課堂解題引導】：
1. 化簡最高心法：『遇事不決，全化正餘弦！』這是三角化簡的第一鐵律。
2. 通分訓練：讓學生練習雙層分數的繁分式化簡。"""
        },
        55: { # 隨堂練習 6
            "title": "🎯 隨堂練習 6 課堂思考方向：",
            "points": [
                "① 倒數關係：cos 35° × sec 35° ＝ 1",
                "② 平方關係：sin² 80° + cos² 80° ＝ 1",
                "③ 平方關係：cot² 12° - csc² 12° ＝ -1"
            ],
            "notes": """【課堂解題引導】：
引導學生熟練套用三大恆等關係，不需算出具體角度的數值即可秒殺！"""
        },
        56: { # 例題 7
            "title": "🎯 例題 7 課堂思考方向（餘角與平方關係化簡）：",
            "points": [
                "① 觀察角度互餘：40°+50°＝90°、72°+18°＝90°",
                "② 餘角代換：tan 50° ＝ cot 40°，得 tan 40° cot 40° ＝ 1",
                "③ 餘角代換：csc 18° ＝ sec 72°，得 tan² 72° - sec² 72° ＝ -1"
            ],
            "notes": """【課堂解題引導】：
1. 破題觀察：角度不是 30°, 45°, 60° 時，立刻觀察兩角是否相加等於 90°！
2. 餘角轉換為同角後，再套用恆等關係化簡。"""
        },
        57: { # 隨堂練習 7
            "title": "🎯 隨堂練習 7 課堂思考方向：",
            "points": [
                "① 互餘配對：35°+55°＝90°，cos 55° ＝ sin 35°",
                "② 平方合一：cos² 35° + sin² 35° ＝ 1",
                "③ 互餘配對：47°+43°＝90°，tan 43° ＝ cot 47°，再套平方關係"
            ],
            "notes": """【課堂解題引導】：
加強互餘角的靈活互換訓練。"""
        },
        58: { # 例題 8
            "title": "🎯 例題 8 課堂思考方向（齊次式同除以 cosθ 技巧）：",
            "points": [
                "① 題目特徵：已知 tan θ 之值，求 sin、cos 的一次齊次式比值",
                "② 秒殺神技：分子分母同除以 cos θ！",
                "③ 全式轉化：所有 sin θ / cos θ 轉化為 tan θ，直接代入已知數值求解"
            ],
            "notes": """【課堂解題引導】：
1. 傳統作法：畫直角三角形求出 sin 和 cos 再代入（很慢且容易帶根號算錯）。
2. 高手作法：同除以 cosθ，直接把全式變成 tanθ 的一次分數式，30 秒內算完！"""
        },
        59: { # 隨堂練習 8
            "title": "🎯 隨堂練習 8 課堂思考方向：",
            "points": [
                "① 觀察為齊次式結構",
                "② 分子分母同除以 cos θ",
                "③ 代入已知條件化簡求值"
            ],
            "notes": """【課堂解題引導】：
齊次式技巧再次演練，深化學生的解題直覺。"""
        },
        64: { # 例題 9
            "title": "🎯 例題 9 課堂思考方向（平方關係與和差乘積轉換）：",
            "points": [
                "① 題目特徵：已知 sin θ + cos θ，求乘積或對稱式",
                "② 必殺技巧：將已知式『兩邊同時平方』！",
                "③ 展開公式：(sin θ + cos θ)² ＝ 1 + 2sin θ cos θ",
                "④ 平方關係代換：解出 sin θ cos θ 後進一步代入求解"
            ],
            "notes": """【課堂解題引導】：
統測極高頻題型！強調 (sinθ ± cosθ)² = 1 ± 2sinθcosθ 的重要性。"""
        },
        65: { # 隨堂練習 9
            "title": "🎯 隨堂練習 9 課堂思考方向：",
            "points": [
                "① 將已知條件兩邊平方求出 sin θ cos θ",
                "② tan θ + cot θ ＝ (sin/cos) + (cos/sin) ＝ 1 / (sin θ cos θ)",
                "③ 帶入乘積倒數求得答案"
            ],
            "notes": """【課堂解題引導】：
引導學生記住常用對稱式 tanθ + cotθ = 1 / (sinθcosθ)。"""
        }
    }

    for idx, info in teaching_data.items():
        slide = prs.slides[idx]
        add_guidance_box(slide, Inches(6.8), Inches(4.3), Inches(5.8), Inches(2.7), info["title"], info["points"])
        set_speaker_notes(slide, info["notes"])
        print(f"Added guidance box & notes to Slide {idx+1}")

    # 針對習題投影片（Slide 68~77，即 idx 67~76）
    exercise_data = {
        67: { # 習題 1
            "title": "💡 習題 1 觀念是非題核心點評：",
            "points": [
                "• (1) ○：一角知六角，畫直角三角形依畢氏定理可定其餘五者",
                "• (2) ○：直角三角形給定兩邊長，銳角三角比完全唯一確定",
                "• (3) ╳：特別角 30° 邊長比 1:√3:2，sin 30° ＝ 1/2",
                "• (4) ○：平方關係 sin²θ + cos²θ ＝ 1 對所有銳角恆成立",
                "• (5) ○：商數關係 tan θ ＝ sin θ / cos θ 為恆等式"
            ],
            "notes": """【習題 1 觀念是非題解析】：
引導學生逐題說明判斷依據，強化定義與定理的掌握。"""
        },
        68: { # 習題 2
            "title": "💡 習題 2 破題關鍵提示：",
            "points": [
                "• 利用畢氏定理求出第三邊長",
                "• 依定義寫出 tan A ＝ 對邊/鄰邊，sec A ＝ 斜邊/鄰邊",
                "• 兩者分母皆為鄰邊，直接分子相加約分求解"
            ],
            "notes": """【習題 2 解析】：
畢氏定理求第三邊後，代入定義相加。"""
        },
        69: { # 習題 3
            "title": "💡 習題 3 破題關鍵提示：",
            "points": [
                "• 設邊長比例參數 k，利用畢氏定理求斜邊",
                "• 依定義寫出 sin A",
                "• 利用餘角關係或互餘三角比求其餘數值"
            ],
            "notes": """【習題 3 解析】：
直角三角形邊長比還原與求值。"""
        },
        70: { # 習題 4
            "title": "💡 習題 4 破題關鍵提示：",
            "points": [
                "• 精確帶入 30°、45°、60° 特別角函數值",
                "• 優先計算平方：sin² 30° ＝ 1/4、cot² 30° ＝ 3、sec² 60° ＝ 4",
                "• 依序通分與乘除化簡"
            ],
            "notes": """【習題 4 解析】：
特別角四則運算精熟。"""
        },
        71: { # 習題 5
            "title": "💡 習題 5 破題關鍵提示：",
            "points": [
                "• 互餘角轉換：20°+70°＝90°，cos 70° ＝ sin 20°",
                "• 平方合一：cos² 20° + sin² 20° ＝ 1",
                "• 互餘轉換：csc 70° ＝ sec 20°，再套平方關係 tan² 20° - sec² 20° ＝ -1"
            ],
            "notes": """【習題 5 解析】：
互餘角與平方關係綜合應用。"""
        },
        72: { # 習題 6
            "title": "💡 習題 6 破題關鍵提示：",
            "points": [
                "• 分式通分合一",
                "• 展開分子利用平方關係 sin²θ + cos²θ ＝ 1 化簡",
                "• 約去公因式得出極簡數值"
            ],
            "notes": """【習題 6 解析】：
通分化簡典範題。"""
        },
        73: { # 習題 7
            "title": "💡 習題 7 破題關鍵提示：",
            "points": [
                "• 觀察式子為齊次式",
                "• 分子分母同除以 cos θ 轉化為 tan θ",
                "• 代入已知條件秒出答案"
            ],
            "notes": """【習題 7 解析】：
齊次式快速法實戰。"""
        },
        74: { # 習題 8
            "title": "💡 習題 8 破題關鍵提示：",
            "points": [
                "• 兩邊平方求出 sin θ cos θ",
                "• tan θ + cot θ 轉化為 1 / (sin θ cos θ)",
                "• 取乘積倒數求得結果"
            ],
            "notes": """【習題 8 解析】：
和差平方轉乘積。"""
        },
        75: { # 習題 9
            "title": "💡 習題 9 破題關鍵提示：",
            "points": [
                "• 進階齊次式化簡技巧",
                "• 分子分母同除以 cos² θ 或因式分解",
                "• 轉化為 tan θ 之二次代數式求值"
            ],
            "notes": """【習題 9 解析】：
進階題二次齊次式處理。"""
        },
        76: { # 習題 10
            "title": "💡 習題 10 破題關鍵提示：",
            "points": [
                "• 兩邊平方解乘積項",
                "• 乘法公式展開：sin³ x + cos³ x ＝ (sin x + cos x)(1 - sin x cos x)",
                "• 逐步代入求值"
            ],
            "notes": """【習題 10 解析】：
立方和公式因式分解。"""
        }
    }

    for idx, info in exercise_data.items():
        slide = prs.slides[idx]
        add_guidance_box(slide, Inches(7.2), Inches(4.5), Inches(5.5), Inches(2.5), info["title"], info["points"])
        set_speaker_notes(slide, info["notes"])
        print(f"Added guidance box & notes to Slide {idx+1} (Exercise)")

    # 3. 插入新頁面
    # 先備知識 4 頁（插入在 Slide 4 目錄之後，即原始 index 4）
    sldIdLst = prs.slides._sldIdLst
    layout_blank = prs.slide_layouts[5] # 課文頁-第一頁
    
    # 建立先備知識 4 頁
    prereq_funcs = [
        create_prerequisite_slide_1,
        create_prerequisite_slide_2,
        create_roadmap_slide,
        create_memory_guide_slide
    ]
    
    # 先加在最後，再移到 index 4
    for i, func in enumerate(prereq_funcs):
        new_s = prs.slides.add_slide(layout_blank)
        for sh in new_s.shapes:
            if sh.is_placeholder:
                sh.text_frame.text = ""
        func(new_s)
        elem = sldIdLst[-1]
        sldIdLst.insert(4 + i, elem)
        print(f"Inserted Prerequisite Slide {i+1} at index {4+i}")

    # 插入超級六邊形（放置在「2-2.3三角函數的恆等關係」重點整理處，原 Slide 44 是重點頁，加上前 4 頁偏移後為 44+4=48）
    # 我們精確找到包含「1. 倒數關係」的投影片位置，將超級六邊形插在它前面
    hex_target_idx = -1
    for idx, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and "1. 倒數關係" in sh.text_frame.text:
                hex_target_idx = idx
                break
        if hex_target_idx != -1:
            break
            
    if hex_target_idx != -1:
        new_hex = prs.slides.add_slide(layout_blank)
        for sh in new_hex.shapes:
            if sh.is_placeholder:
                sh.text_frame.text = ""
        create_super_hexagon_slide(new_hex)
        hex_elem = sldIdLst[-1]
        sldIdLst.insert(hex_target_idx, hex_elem)
        print(f"Inserted Super Hexagon at index {hex_target_idx}")

    # 存檔
    print("Saving modified presentation...")
    prs.save(OUT_FILE)
    prs.save(COPY_FILE)
    print(f"Successfully modified and saved to {OUT_FILE} and {COPY_FILE}!")

if __name__ == "__main__":
    run_modifications()
