# -*- coding: utf-8 -*-
"""
generate_landscape_worksheets.py
泰宇版 技術高中數學（第一冊）2-2 三角函數的定義與性質
JIS B4 橫向雙欄 Word 學習單產生器 (學生自主學習單 + 教師教學評量解答版)
遵循《SOP v3.0 名偵探柯南數理探案特訓 ‧ 旗艦全功能標準版》規範：
- Word 原生 OMML 數學方程式 (<m:oMath>)
- JIS B4 橫向 (364mm x 257mm) 雙欄排版 (含中央垂直分隔線)
- 學生版留白充足計算區，教師版醒目紅色 (#DC2626) 原生方程式解答
- 自動將產出之 docx 編碼為 base64 寫入 docx_b64.py
"""

import os
import re
import base64
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree
import latex2mathml.converter

# 尋找微軟 Office 內建之 MML2OMML.XSL
XSL_PATHS = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
]
XSL_FILE = next((p for p in XSL_PATHS if os.path.exists(p)), None)
if XSL_FILE:
    XSLT_DOC = etree.parse(XSL_FILE)
    TRANSFORM = etree.XSLT(XSLT_DOC)
    print(f"成功載入 MML2OMML.XSL: {XSL_FILE}")
else:
    TRANSFORM = None
    print("警告：未找到 MML2OMML.XSL，將以純文字替代！")

def latex_to_omml(latex_str, is_teacher=False):
    """將 LaTeX 數學式轉換為 Word 原生 OMML (<m:oMath>) XML 節點"""
    if not TRANSFORM:
        return None
    try:
        clean_latex = latex_str.replace(r'\dfrac', r'\frac')
        clean_latex = clean_latex.replace(r'\,', r' ')
        clean_latex = clean_latex.replace(r'\approx', r'~')
        clean_latex = clean_latex.replace(r'\cdot', r'\times')
        
        mathml_str = latex2mathml.converter.convert(clean_latex)
        mml_tree = etree.fromstring(mathml_str.encode('utf-8'))
        omml_tree = TRANSFORM(mml_tree)
        omml_xml = etree.tostring(omml_tree, encoding='utf-8').decode('utf-8')
        omml_xml = re.sub(r'<\?xml[^>]+\?>', '', omml_xml).strip()
        
        if is_teacher:
            omml_xml = omml_xml.replace(
                '<m:oMath ',
                '<m:oMath xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            )
            omml_xml = re.sub(
                r'(<m:r>|<m:r\s+[^>]*>)',
                r'\1<m:rPr><w:rPr><w:color w:val="DC2626"/></w:rPr></m:rPr>',
                omml_xml
            )
        return parse_xml(omml_xml)
    except Exception as e:
        return None

def add_math_paragraph(paragraph, full_text, is_teacher=False, default_color=None, bold=False):
    """混合純文字與 $...$ 數學式，自動切分並嵌入原生 OMML 節點"""
    parts = re.split(r'(\$[^\$]+\$)', full_text)
    for part in parts:
        if not part:
            continue
        if part.startswith('$') and part.endswith('$') and len(part) > 2:
            latex_content = part[1:-1].strip()
            omml_node = latex_to_omml(latex_content, is_teacher=is_teacher)
            if omml_node is not None:
                paragraph._p.append(omml_node)
            else:
                run = paragraph.add_run(latex_content)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.bold = bold
                if is_teacher:
                    run.font.color.rgb = RGBColor(220, 38, 38)
                elif default_color:
                    run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(part)
            run.font.name = "微軟正黑體"
            run.font.size = Pt(11)
            run.bold = bold
            if is_teacher:
                run.font.color.rgb = RGBColor(220, 38, 38)
            elif default_color:
                run.font.color.rgb = default_color

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """設定儲存格邊距 (dxa)"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_box(doc, text, color_rgb=(180, 83, 9)):
    """加入精緻標題區塊"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"【{text}】")
    run.font.name = "微軟正黑體"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(*color_rgb)

def add_box(doc, title, content_lines, is_teacher=False, answer_lines=None, blank_lines=4):
    """加入題目與計算方框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="D97706"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_math_paragraph(p, title, default_color=RGBColor(120, 53, 15), bold=True)
    
    for line in content_lines:
        lp = cell.add_paragraph()
        lp.paragraph_format.space_after = Pt(2)
        add_math_paragraph(lp, line)
        
    if is_teacher and answer_lines:
        ap = cell.add_paragraph()
        ap.paragraph_format.space_before = Pt(4)
        ap.paragraph_format.space_after = Pt(2)
        arun = ap.add_run("【👓 名偵探破案解析 ‧ CASE CLOSED】：")
        arun.font.name = "微軟正黑體"
        arun.font.size = Pt(10.5)
        arun.bold = True
        arun.font.color.rgb = RGBColor(220, 38, 38)
        for aline in answer_lines:
            alp = cell.add_paragraph()
            alp.paragraph_format.space_after = Pt(2)
            add_math_paragraph(alp, aline, is_teacher=True)
    else:
        for _ in range(blank_lines):
            sp = cell.add_paragraph()
            sp.paragraph_format.space_after = Pt(0)
            srun = sp.add_run("　")
            srun.font.size = Pt(11)

def create_worksheet(is_teacher=False):
    doc = docx.Document()
    
    # 1. 版面設定：JIS B4 橫向 (364mm x 257mm)
    section = doc.sections[0]
    section.page_width = Mm(364)
    section.page_height = Mm(257)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    
    # 2. 雙欄配置 (含垂直分隔線)
    sectPr = section._sectPr
    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:sep'), '1')
    cols.set(qn('w:space'), '720')
    sectPr.append(cols)
    
    # 3. 頁尾
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    status_str = "[警視廳機密解答卷宗 ‧ 密碼 88888888]" if is_teacher else "[偵探搜查實作單]"
    f_run = f_p.add_run(f"泰宇版 技術高中數學（第一冊）2-2 三角函數的定義與性質 {status_str} ｜ PLVS 數學科  教育學習用")
    f_run.font.name = "微軟正黑體"
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = RGBColor(140, 140, 140)
    
    # 4. 標題橫幅
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title_p.add_run("泰宇版 技術高中數學（第一冊）2-2 三角函數的定義與性質 數位學習單")
    t_run.font.name = "微軟正黑體"
    t_run.font.size = Pt(16)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(120, 53, 15)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mode_text = "【名偵探柯南數理探案 ‧ 警視廳機密解答卷宗（紅字 OMML 詳解）】" if is_teacher else "【名偵探柯南數理探案 ‧ 偵探搜查實作單（鷹架引導自主學習）】"
    s_run = sub_p.add_run(mode_text)
    s_run.font.name = "微軟正黑體"
    s_run.font.size = Pt(11)
    s_run.bold = True
    s_run.font.color.rgb = RGBColor(220, 38, 38) if is_teacher else RGBColor(180, 83, 9)
    
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_p.add_run("班級：___________　座號：___________　姓名：_______________　特訓評量：__________ / 100")
    info_run.font.name = "微軟正黑體"
    info_run.font.size = Pt(10)
    info_run.font.color.rgb = RGBColor(71, 85, 105)
    
    # ==========================================
    # 模組一：核心觀念奠基與定理填空
    # ==========================================
    add_heading_box(doc, "模組一：【名偵探破案密碼】核心觀念奠基與定理填空", (22, 101, 52))
    
    add_box(
        doc,
        "重點 1：銳角三角函數的六大嚴格定義",
        [
            "在直角 $\\triangle$ 中，角 $\\theta$ 的斜邊為 $r$、對邊為 $y$、鄰邊為 $x$：",
            "1. 正弦 $\\sin\\theta = $___________ ； 餘割 $\\csc\\theta = $___________",
            "2. 餘弦 $\\cos\\theta = $___________ ； 正割 $\\sec\\theta = $___________",
            "3. 正切 $\\tan\\theta = $___________ ； 餘切 $\\cot\\theta = $___________",
            "記憶口訣：S 配 C，T 配 T；正弦配餘割，餘弦配正割，正切配餘切！"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. $\\sin\\theta = \\frac{y}{r} = \\frac{\\text{對邊}}{\\text{斜邊}}$ ； $\\csc\\theta = \\frac{r}{y} = \\frac{\\text{斜邊}}{\\text{對邊}}$",
            "2. $\\cos\\theta = \\frac{x}{r} = \\frac{\\text{鄰邊}}{\\text{斜邊}}$ ； $\\sec\\theta = \\frac{r}{x} = \\frac{\\text{斜邊}}{\\text{鄰邊}}$",
            "3. $\\tan\\theta = \\frac{y}{x} = \\frac{\\text{對邊}}{\\text{鄰邊}}$ ； $\\cot\\theta = \\frac{x}{y} = \\frac{\\text{鄰邊}}{\\text{對邊}}$"
        ] if is_teacher else None,
        blank_lines=2
    )
    
    add_box(
        doc,
        "重點 2：特別角 $30^\\circ$、$45^\\circ$、$60^\\circ$ 精確值速記表",
        [
            "1. $30^\\circ = \\frac{\\pi}{6}$：$\\sin 30^\\circ = $___________，$\\cos 30^\\circ = $___________，$\\tan 30^\\circ = $___________",
            "2. $45^\\circ = \\frac{\\pi}{4}$：$\\sin 45^\\circ = $___________，$\\cos 45^\\circ = $___________，$\\tan 45^\\circ = $___________",
            "3. $60^\\circ = \\frac{\\pi}{3}$：$\\sin 60^\\circ = $___________，$\\cos 60^\\circ = $___________，$\\tan 60^\\circ = $___________"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. $\\sin 30^\\circ = \\frac{1}{2}$，$\\cos 30^\\circ = \\frac{\\sqrt{3}}{2}$，$\\tan 30^\\circ = \\frac{\\sqrt{3}}{3}$",
            "2. $\\sin 45^\\circ = \\frac{\\sqrt{2}}{2}$，$\\cos 45^\\circ = \\frac{\\sqrt{2}}{2}$，$\\tan 45^\\circ = 1$",
            "3. $\\sin 60^\\circ = \\frac{\\sqrt{3}}{2}$，$\\cos 60^\\circ = \\frac{1}{2}$，$\\tan 60^\\circ = \\sqrt{3}$"
        ] if is_teacher else None,
        blank_lines=2
    )
    
    add_box(
        doc,
        "重點 3：神奇六角星陣四大恆等關係",
        [
            "1. 倒數關係：$\\sin\\theta \\cdot \\csc\\theta = $______ ； $\\cos\\theta \\cdot \\sec\\theta = $______ ； $\\tan\\theta \\cdot \\cot\\theta = $______",
            "2. 商數關係：$\\tan\\theta = $___________ ； $\\cot\\theta = $___________",
            "3. 平方關係：$\\sin^2\\theta + \\cos^2\\theta = $______ ； $1 + \\tan^2\\theta = $______ ； $1 + \\cot^2\\theta = $______",
            "4. 餘角關係：$\\sin(90^\\circ-\\theta) = $___________ ； $\\tan(90^\\circ-\\theta) = $___________ ； $\\sec(90^\\circ-\\theta) = $___________"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. 倒數乘積皆為 1。",
            "2. $\\tan\\theta = \\frac{\\sin\\theta}{\\cos\\theta}$ ； $\\cot\\theta = \\frac{\\cos\\theta}{\\sin\\theta}$",
            "3. $\\sin^2\\theta + \\cos^2\\theta = 1$ ； $1 + \\tan^2\\theta = \\sec^2\\theta$ ； $1 + \\cot^2\\theta = \\csc^2\\theta$",
            "4. 橫向端點換裝：$\\sin(90^\\circ-\\theta)=\\cos\\theta$ ； $\\tan(90^\\circ-\\theta)=\\cot\\theta$ ； $\\sec(90^\\circ-\\theta)=\\csc\\theta$"
        ] if is_teacher else None,
        blank_lines=2
    )

    add_box(
        doc,
        "重點 4：常用三角恆等式與生活素養",
        [
            "1. $(\\sin\\theta \\pm \\cos\\theta)^2 = 1 \\pm $___________",
            "2. $\\tan\\theta + \\cot\\theta = $___________",
            "3. 生活坡度計算：坡度 $= \\tan\\theta = \\frac{\\text{垂直上升}}{\\text{水平前進}}$。已知正切比值，按計算機___________鍵反推求角度。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. $(\\sin\\theta \\pm \\cos\\theta)^2 = 1 \\pm 2\\sin\\theta\\cos\\theta$",
            "2. $\\tan\\theta + \\cot\\theta = \\frac{1}{\\sin\\theta\\cos\\theta}$",
            "3. 按反三角函數鍵 $\\tan^{-1}$（或 SHIFT + tan）。"
        ] if is_teacher else None,
        blank_lines=2
    )

    # ==========================================
    # 模組二：課本核心範例與隨堂成對實戰演練
    # ==========================================
    add_heading_box(doc, "模組二：課本核心範例與隨堂成對實戰演練", (30, 58, 138))
    
    # 範例 1 & 隨堂 1
    add_box(
        doc,
        "【CASE 04 & 06】範例 1(1) 與 隨堂 1(1)：8-15-17 直角三角形雙重視角演練",
        [
            "題目：在直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ, AC = 8, BC = 15, AB = 17$。",
            "(1) 求 $\\angle A$ 的六個三角函數值：",
            "(2) 求 $\\angle B$ 的六個三角函數值："
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\angle A$ 視角：對邊 15、鄰邊 8、斜邊 17。$\\sin A = \\frac{15}{17}, \\cos A = \\frac{8}{17}, \\tan A = \\frac{15}{8}, \\csc A = \\frac{17}{15}, \\sec A = \\frac{17}{8}, \\cot A = \\frac{8}{15}$。",
            "(2) $\\angle B$ 視角：對邊 8、鄰邊 15、斜邊 17。$\\sin B = \\frac{8}{17}, \\cos B = \\frac{15}{17}, \\tan B = \\frac{8}{15}, \\csc B = \\frac{17}{8}, \\sec B = \\frac{17}{15}, \\cot B = \\frac{15}{8}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 2 & 隨堂 2
    add_box(
        doc,
        "【CASE 08 & 09】範例 2 與 隨堂 2：知一三角函數值求其餘五者",
        [
            "(1) [範例 2] 直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ$ 且 $\\cos A = \\frac{5}{13}$，求 $\\angle A$ 其餘各三角函數值。",
            "(2) [隨堂 2] 直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ$ 且 $\\tan A = \\frac{24}{7}$，求 $\\angle A$ 其餘各三角函數值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 設鄰邊 $5k$、斜邊 $13k$，對邊 $BC = \\sqrt{13^2 - 5^2}k = 12k$。$\\sin A = \\frac{12}{13}, \\tan A = \\frac{12}{5}, \\csc A = \\frac{13}{12}, \\sec A = \\frac{13}{5}, \\cot A = \\frac{5}{12}$。",
            "(2) 設對邊 $24k$、鄰邊 $7k$，斜邊 $AB = \\sqrt{24^2 + 7^2}k = 25k$。$\\sin A = \\frac{24}{25}, \\cos A = \\frac{7}{25}, \\csc A = \\frac{25}{24}, \\sec A = \\frac{25}{7}, \\cot A = \\frac{7}{24}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 3 & 隨堂 3
    add_box(
        doc,
        "【CASE 10 & 11】範例 3 與 隨堂 3：比例關係求實際邊長與函數值",
        [
            "(1) [範例 3] 直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ, \\sin A = \\frac{4}{5}, AB = 15$。求 $\\cos A$、$\\tan A$、$BC$ 與 $AC$ 之長。",
            "(2) [隨堂 3] 直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ, \\csc A = \\frac{17}{15}, AC = 40$。求 $\\tan A$、$\\sec A$、$AB$ 與 $BC$ 之長。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 斜邊 $5k = 15 \\implies k = 3$。$\\cos A = \\frac{3}{5}, \\tan A = \\frac{4}{3}, BC = 4(3) = 12, AC = 3(3) = 9$。",
            "(2) $\\sin A = \\frac{15}{17}$，三邊比 $8k : 15k : 17k$。$AC = 8k = 40 \\implies k = 5$。$\\tan A = \\frac{15}{8}, \\sec A = \\frac{17}{8}, AB = 17(5) = 85, BC = 15(5) = 75$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 4 & 隨堂 4
    add_box(
        doc,
        "【CASE 15 ~ 18】範例 4 與 隨堂 4：特別角與弳度量混合求值",
        [
            "(1) 試求 $\\sin 30^\\circ - \\cos 60^\\circ + \\tan 45^\\circ$ 之值。",
            "(2) 試求 $2\\sin^2 45^\\circ + \\cot 30^\\circ \\csc 60^\\circ$ 之值。",
            "(3) 試求 $\\sin^2 \\frac{\\pi}{4} + \\cos^2 \\frac{\\pi}{6} - \\cot^2 \\frac{\\pi}{3}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 原式 $= \\frac{1}{2} - \\frac{1}{2} + 1 = 1$。",
            "(2) 原式 $= 2(\\frac{1}{2}) + \\sqrt{3}(\\frac{2}{\\sqrt{3}}) = 1 + 2 = 3$。",
            "(3) 原式 $= \\frac{1}{2} + \\frac{3}{4} - \\frac{1}{3} = \\frac{6 + 9 - 4}{12} = \\frac{11}{12}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 6 & 隨堂 6
    add_box(
        doc,
        "【CASE 26 ~ 31】範例 6 與 隨堂 6：六角星陣恆等關係基礎求值",
        [
            "(1) 求 $\\sin 50^\\circ \\times \\csc 50^\\circ$ 之值。",
            "(2) 求 $\\sec^2 70^\\circ - \\tan^2 70^\\circ$ 之值。",
            "(3) 求 $\\cot^2 12^\\circ - \\csc^2 12^\\circ$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 倒數關係：原式 $= 1$。",
            "(2) 平方關係：$1 + \\tan^2\\theta = \\sec^2\\theta \\implies \\sec^2\\theta - \\tan^2\\theta = 1$。",
            "(3) 平方關係：$1 + \\cot^2\\theta = \\csc^2\\theta \\implies \\cot^2\\theta - \\csc^2\\theta = -1$。"
        ] if is_teacher else None,
        blank_lines=3
    )

    # 範例 7 & 隨堂 7
    add_box(
        doc,
        "【CASE 32 & 33】範例 7 與 隨堂 7：餘角轉化與平方綜合求值",
        [
            "(1) 試求 $\\tan 40^\\circ \\tan 50^\\circ + \\tan^2 72^\\circ - \\csc^2 18^\\circ$ 之值。",
            "(2) 試求 $\\cos^2 35^\\circ + \\cos^2 55^\\circ + \\csc^2 47^\\circ - \\tan^2 43^\\circ$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\tan 50^\\circ = \\cot 40^\\circ \\implies \\tan 40^\\circ\\cot 40^\\circ = 1$；$\\csc 18^\\circ = \\sec 72^\\circ \\implies \\tan^2 72^\\circ - \\sec^2 72^\\circ = -1$。原式 $= 1 + (-1) = 0$。",
            "(2) $\\cos 55^\\circ = \\sin 35^\\circ \\implies \\cos^2 35^\\circ + \\sin^2 35^\\circ = 1$；$\\tan 43^\\circ = \\cot 47^\\circ \\implies \\csc^2 47^\\circ - \\cot^2 47^\\circ = 1$。原式 $= 1 + 1 = 2$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 8 & 隨堂 8
    add_box(
        doc,
        "【CASE 34 & 35】範例 8 與 隨堂 8：一次齊次式商數轉化神技",
        [
            "(1) 已知 $\\theta$ 為銳角且 $\\tan\\theta = 2$，求 $\\frac{\\sin\\theta + 2\\cos\\theta}{3\\sin\\theta - \\cos\\theta}$ 之值。",
            "(2) 已知 $\\theta$ 為銳角且 $\\cot\\theta = \\frac{2}{3}$，求 $\\frac{\\sin\\theta - 3\\cos\\theta}{5\\sin\\theta + 6\\cos\\theta}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 分子分母同除以 $\\cos\\theta$ 得 $\\frac{\\tan\\theta + 2}{3\\tan\\theta - 1} = \\frac{2 + 2}{3(2) - 1} = \\frac{4}{5}$。",
            "(2) 分子分母同除以 $\\sin\\theta$ 得 $\\frac{1 - 3\\cot\\theta}{5 + 6\\cot\\theta} = \\frac{1 - 3(\\frac{2}{3})}{5 + 6(\\frac{2}{3})} = \\frac{1 - 2}{5 + 4} = -\\frac{1}{9}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 範例 9 & 隨堂 9
    add_box(
        doc,
        "【CASE 37 ~ 40】範例 9 與 隨堂 9：和差連環扣與正餘切之和",
        [
            "(1) 已知銳角 $\\theta$ 滿足 $\\sin\\theta + \\cos\\theta = \\frac{4}{3}$，試求 $\\sin\\theta\\cos\\theta$ 與 $\\tan\\theta + \\cot\\theta$ 之值。",
            "(2) 已知銳角 $\\theta$ 滿足 $\\sin\\theta - \\cos\\theta = \\frac{3}{5}$，試求 $\\sin\\theta\\cos\\theta$ 與 $\\tan\\theta + \\cot\\theta$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 兩邊平方：$1 + 2\\sin\\theta\\cos\\theta = \\frac{16}{9} \\implies \\sin\\theta\\cos\\theta = \\frac{7}{18}$；$\\tan\\theta + \\cot\\theta = \\frac{1}{\\sin\\theta\\cos\\theta} = \\frac{18}{7}$。",
            "(2) 兩邊平方：$1 - 2\\sin\\theta\\cos\\theta = \\frac{9}{25} \\implies \\sin\\theta\\cos\\theta = \\frac{8}{25}$；$\\tan\\theta + \\cot\\theta = \\frac{1}{\\sin\\theta\\cos\\theta} = \\frac{25}{8}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # ==========================================
    # 模組三：課後習題 ‧ 觀念釐清是非專題
    # ==========================================
    add_heading_box(doc, "模組三：課後習題 ‧ 觀念釐清是非專題（是非題）", (194, 65, 12))
    
    add_box(
        doc,
        "【CASE 41 ~ 45】課本 P.123 觀念釐清區（是非題五大案判定）",
        [
            "(　)(1) 已知 $\\theta$ 為銳角，若已知 $\\sin\\theta$，則可求出剩餘的五個三角函數值。",
            "(　)(2) 直角三角形 $ABC$ 中，$\\angle C = 90^\\circ$，若 $\\sin A = \\frac{3}{5}$，則可求出 $AB$、$BC$、$CA$ 的長度。",
            "(　)(3) 若 $\\alpha$、$\\beta$ 為銳角，由平方關係可得 $\\sin^2\\alpha + \\cos^2\\beta = 1$。",
            "(　)(4) 若 $\\theta$ 為銳角，則 $\\sin\\theta \\times \\cos\\theta = 1$。",
            "(　)(5) 若 $\\theta$ 為銳角，則 $(\\sin\\theta + \\cos\\theta)^2 = \\sin^2\\theta + \\cos^2\\theta = 1$。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(○)(1) 正確！已知銳角三角比，利用畢氏定理可定三邊比例，其餘五個函數值唯一確定。",
            "(×)(2) 錯誤！三角函數是比值而非實際長度，無邊長數據無法得知實際長度。",
            "(×)(3) 錯誤！平方關係嚴格限定同角，角度不同不保證等於 1。",
            "(×)(4) 錯誤！倒數關係為 $\\sin\\theta \\times \\csc\\theta = 1$。",
            "(×)(5) 錯誤！遺漏交叉項：$(\\sin\\theta+\\cos\\theta)^2 = 1 + 2\\sin\\theta\\cos\\theta > 1$。"
        ] if is_teacher else None,
        blank_lines=2
    )

    # ==========================================
    # 模組四：課後習題 ‧ 基礎進階挑戰與破案筆記
    # ==========================================
    add_heading_box(doc, "模組四：課後習題 ‧ 基礎進階挑戰與破案筆記", (88, 28, 135))
    
    add_box(
        doc,
        "【CASE 46 & 51】習題基礎題精選實戰",
        [
            "(1) 直角 $\\triangle ABC$ 中，$\\angle C = 90^\\circ, \\sin A = \\frac{3}{5}$，求 $\\tan A + \\sec A$ 之值。",
            "(2) 試求 $\\cos^2 20^\\circ + \\cos^2 70^\\circ + \\tan^2 20^\\circ - \\csc^2 70^\\circ$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 三邊為 3-4-5，$\\tan A = \\frac{3}{4}, \\sec A = \\frac{5}{4} \\implies \\frac{3+5}{4} = 2$。",
            "(2) $\\cos^2 20^\\circ + \\sin^2 20^\\circ = 1$ ； $\\tan^2 20^\\circ - \\sec^2 20^\\circ = -1$。原式 $= 1 + (-1) = 0$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    add_box(
        doc,
        "【CASE 56 & 59】習題進階題 ‧ 劇場版終極破案",
        [
            "(1) 已知 $\\theta$ 為銳角且 $\\cos\\theta = \\frac{1}{5}$，試求 $\\frac{1}{1 + \\sin\\theta} + \\frac{1}{1 - \\sin\\theta}$ 之值。",
            "(2) 已知 $x$ 為銳角且 $\\tan x + \\cot x = \\frac{25}{12}$，求 $\\sin^3 x + \\cos^3 x$ 之值。（提示：$a^3 + b^3 = (a+b)(a^2 - ab + b^2)$）"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 通分得 $\\frac{2}{1 - \\sin^2\\theta} = \\frac{2}{\\cos^2\\theta} = \\frac{2}{(\\frac{1}{5})^2} = 2 \\times 25 = 50$。",
            "(2) $\\sin x\\cos x = \\frac{12}{25}$ ； $(\\sin x+\\cos x)^2 = 1 + 2(\\frac{12}{25}) = \\frac{49}{25} \\implies \\sin x+\\cos x = \\frac{7}{5}$。<br>由立方和公式：$(\\frac{7}{5})(1 - \\frac{12}{25}) = \\frac{7}{5} \\times \\frac{13}{25} = \\frac{91}{125}$。"
        ] if is_teacher else None,
        blank_lines=4
    )

    # 偵探搜查反思筆記
    add_box(
        doc,
        "【偵探搜查實作反思與心得筆記】",
        [
            "1. 本單元我最拿手的主題是：___________________________________________________________",
            "2. 我曾踩過或特別需要注意的陷阱是：___________________________________________________",
            "3. 少年偵探團升級特訓考核達成關卡：STAGE [  ]  /  總破案積分：__________ 分"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "教師評語與回饋：本單元學生已完整建立銳角三角函數概念，能熟練運用六角形恆等式進行化簡與解題。"
        ] if is_teacher else None,
        blank_lines=2
    )

    # 儲存檔案
    out_name = "2-2_三角函數的定義與性質_教師解答版.docx" if is_teacher else "2-2_三角函數的定義與性質_學生學習單.docx"
    doc.save(out_name)
    print(f"成功產出: {out_name}")
    return out_name

if __name__ == "__main__":
    student_file = create_worksheet(is_teacher=False)
    teacher_file = create_worksheet(is_teacher=True)
    
    # 產出 Base64 字典檔案供前端網頁直接下載
    with open(student_file, "rb") as f:
        s_b64 = base64.b64encode(f.read()).decode("utf-8")
    with open(teacher_file, "rb") as f:
        t_b64 = base64.b64encode(f.read()).decode("utf-8")
        
    with open("docx_b64.py", "w", encoding="utf-8") as f:
        f.write("# -*- coding: utf-8 -*-\n")
        f.write('STUDENT_DOCX_B64 = "' + s_b64 + '"\n')
        f.write('TEACHER_DOCX_B64 = "' + t_b64 + '"\n')
    print("成功生成 docx_b64.py！")
